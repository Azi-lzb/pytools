from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from pytools.core.config_init import initialize_or_repair_config
from pytools.core.config_xlsx import SHEET_FILE_SORT_CONFIG
from pytools.core.convert_33_34_35_37 import run_file_sort_prefix_49


def _make_cfg(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_FILE_SORT_CONFIG
    ws.append(["是否启用", "排序前缀", "匹配关键词", "备注"])
    ws.append(["Y", "1", "政策性银行;国开行;农发行", "政策性银行"])
    ws.append(["Y", "02", "工商银行;农业银行;中国银行;建设银行;交通银行", "国有银行"])
    ws.append(["Y", "003", "招商银行;浦发银行;中信银行", "股份制银行"])
    ws.append(["N", "004", "不应命中", "禁用"])
    wb.save(path)
    wb.close()


def test_file_sort_prefix_adds_three_digit_prefix_and_uses_first_match(tmp_path):
    cfg = tmp_path / "config.xlsx"
    out_dir = tmp_path / "output"
    logs = tmp_path / "logs"
    _make_cfg(cfg)
    src1 = tmp_path / "202506_建设银行_存款表.xlsx"
    src2 = tmp_path / "202506_国开行_贷款表.xlsx"
    src3 = tmp_path / "202506_未命中文件.xlsx"
    src1.write_text("x", encoding="utf-8")
    src2.write_text("x", encoding="utf-8")
    src3.write_text("x", encoding="utf-8")

    stat = run_file_sort_prefix_49(cfg, [src1, src2, src3], out_dir, logs)

    assert stat["files"] == 3
    assert stat["renamed"] == 2
    assert stat["unmatched"] == 1
    assert (tmp_path / "002_202506_建设银行_存款表.xlsx").exists()
    assert (tmp_path / "001_202506_国开行_贷款表.xlsx").exists()
    assert src3.exists()

    result = pd.read_excel(stat["output"], sheet_name="文件排序前缀结果", dtype=object).fillna("")
    assert list(result["排序前缀"])[:3] == ["002", "001", ""]
    assert list(result["状态"])[:3] == ["已重命名", "已重命名", "未命中"]


def test_file_sort_prefix_replaces_existing_prefix_and_uses_next_path(tmp_path):
    cfg = tmp_path / "config.xlsx"
    out_dir = tmp_path / "output"
    logs = tmp_path / "logs"
    _make_cfg(cfg)
    src = tmp_path / "01_202506_建设银行_存款表.xlsx"
    existing = tmp_path / "002_202506_建设银行_存款表.xlsx"
    src.write_text("x", encoding="utf-8")
    existing.write_text("old", encoding="utf-8")

    stat = run_file_sort_prefix_49(cfg, [src], out_dir, logs)

    assert stat["renamed"] == 1
    assert existing.exists()
    assert (tmp_path / "002_202506_建设银行_存款表_1.xlsx").exists()
    result = pd.read_excel(stat["output"], sheet_name="文件排序前缀结果", dtype=object).fillna("")
    assert result.loc[0, "状态"] == "目标同名：已自动后缀"
    assert result.loc[0, "新文件名"] == "002_202506_建设银行_存款表_1.xlsx"


def test_file_sort_prefix_replaces_dash_and_space_prefixes(tmp_path):
    cfg = tmp_path / "config.xlsx"
    out_dir = tmp_path / "output"
    logs = tmp_path / "logs"
    _make_cfg(cfg)
    src1 = tmp_path / "02-202506_国开行_贷款表.xlsx"
    src2 = tmp_path / "003 202506_浦发银行_存款表.xlsx"
    src1.write_text("x", encoding="utf-8")
    src2.write_text("x", encoding="utf-8")

    run_file_sort_prefix_49(cfg, [src1, src2], out_dir, logs)

    assert (tmp_path / "001_202506_国开行_贷款表.xlsx").exists()
    assert (tmp_path / "003_202506_浦发银行_存款表.xlsx").exists()


def test_file_sort_prefix_uses_longest_enabled_prefix_width(tmp_path):
    cfg = tmp_path / "config.xlsx"
    out_dir = tmp_path / "output"
    logs = tmp_path / "logs"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_FILE_SORT_CONFIG
    ws.append(["是否启用", "排序前缀", "匹配关键词", "备注"])
    ws.append(["Y", "1", "建设银行", "一位输入"])
    ws.append(["Y", "02", "招商银行", "两位输入"])
    ws.append(["N", "003", "禁用银行", "禁用三位不参与宽度判断"])
    wb.save(cfg)
    wb.close()
    src1 = tmp_path / "202506_建设银行_存款表.xlsx"
    src2 = tmp_path / "202506_招商银行_存款表.xlsx"
    src1.write_text("x", encoding="utf-8")
    src2.write_text("x", encoding="utf-8")

    run_file_sort_prefix_49(cfg, [src1, src2], out_dir, logs)

    assert (tmp_path / "01_202506_建设银行_存款表.xlsx").exists()
    assert (tmp_path / "02_202506_招商银行_存款表.xlsx").exists()


def test_config_init_creates_file_sort_config(tmp_path):
    cfg = tmp_path / "config.xlsx"

    initialize_or_repair_config(cfg)

    wb = load_workbook(cfg)
    try:
        ws = wb[SHEET_FILE_SORT_CONFIG]
        headers = [ws.cell(1, c).value for c in range(1, 5)]
        assert headers == ["是否启用", "排序前缀", "匹配关键词", "备注"]
        assert ws["A2"].value == "Y"
        assert ws["B2"].value == "001"
    finally:
        wb.close()
