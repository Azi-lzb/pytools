from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from pytools.core.config_xlsx import ARCHIVE_TYPE_CONFIG_COLS, SHEET_ARCHIVE_TYPE_CONFIG, validate_sheets_for_feature
from pytools.core.config_init import initialize_or_repair_config
from pytools.core.convert_33_34_35_37 import (
    run_archive_by_date_copy,
    run_archive_by_date_preview,
    run_archive_plan_copy,
    run_archive_by_type_copy,
    run_archive_by_type_preview,
)


def _touch(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("x", encoding="utf-8")


def _make_type_cfg(path: Path, include_exts: str = "", exclude_exts: str = "") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_ARCHIVE_TYPE_CONFIG
    for col, name in enumerate(ARCHIVE_TYPE_CONFIG_COLS, start=1):
        ws.cell(1, col).value = name
    rows = [
        ("是", "普惠贷款表", "普惠贷款表", ""),
        ("是", "科技贷款表", "科技贷款表", ""),
        ("是", "贷款表", "贷款表", ""),
        ("否", "禁用表", "禁用表", ""),
    ]
    for row_no, row in enumerate(rows, start=2):
        for col_no, value in enumerate(row, start=1):
            ws.cell(row_no, col_no).value = value
    ws["E2"].value = include_exts
    ws["F2"].value = exclude_exts
    wb.save(path)
    wb.close()


def _read_plan(path: Path) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name="归档计划", dtype=str).fillna("")


def test_archive_by_date_detects_formats_and_unrecognized(tmp_path):
    src = tmp_path / "src"
    _touch(src / "20260609_存款表.xlsx")
    _touch(src / "2026-06-09_贷款表.xlsx")
    _touch(src / "2026年6月普惠贷款表.xlsx")
    _touch(src / "2025.03.31_存款表.xlsx")
    _touch(src / "20250331广东省信贷收支月报.xls")
    _touch(src / "无日期_科技贷款表.xlsx")

    stat = run_archive_by_date_preview([src], tmp_path / "archive", tmp_path / "out", tmp_path / "logs")
    df = _read_plan(stat["output"])

    by_name = {row["文件名"]: row for _, row in df.iterrows()}
    assert by_name["20260609_存款表.xlsx"]["分类结果"] == "202606"
    assert by_name["2026-06-09_贷款表.xlsx"]["分类结果"] == "202606"
    assert by_name["2026年6月普惠贷款表.xlsx"]["分类结果"] == "202606"
    assert by_name["2025.03.31_存款表.xlsx"]["分类结果"] == "202503"
    assert by_name["20250331广东省信贷收支月报.xls"]["分类结果"] == "202503"
    assert by_name["无日期_科技贷款表.xlsx"]["分类结果"] == "未识别日期"
    assert by_name["无日期_科技贷款表.xlsx"]["状态"] == "未识别日期"


def test_archive_by_type_uses_config_order_and_fallback(tmp_path):
    cfg = tmp_path / "config.xlsx"
    _make_type_cfg(cfg)
    src = tmp_path / "src"
    _touch(src / "202606_普惠贷款表.xlsx")
    _touch(src / "202606_贷款表.xlsx")
    _touch(src / "202606_禁用表.xlsx")
    _touch(src / "nested" / "202606_科技贷款表.xlsx")

    stat = run_archive_by_type_preview(cfg, [src], tmp_path / "archive", tmp_path / "out", tmp_path / "logs")
    df = _read_plan(stat["output"])

    by_name = {row["文件名"]: row for _, row in df.iterrows()}
    assert stat["files"] == 4
    assert stat["rules"] == 3
    assert by_name["202606_普惠贷款表.xlsx"]["分类结果"] == "普惠贷款表"
    assert by_name["202606_贷款表.xlsx"]["分类结果"] == "贷款表"
    assert by_name["202606_科技贷款表.xlsx"]["分类结果"] == "科技贷款表"
    assert by_name["202606_禁用表.xlsx"]["分类结果"] == "未识别类型"
    assert by_name["202606_禁用表.xlsx"]["状态"] == "未识别类型"


def test_archive_preview_uses_suffix_for_existing_target(tmp_path):
    src = tmp_path / "src"
    target = tmp_path / "archive"
    _touch(src / "202606_存款表.xlsx")
    _touch(target / "202606" / "202606_存款表.xlsx")

    stat = run_archive_by_date_preview([src], target, tmp_path / "out", tmp_path / "logs")
    df = _read_plan(stat["output"])

    row = df.iloc[0]
    assert row["目标文件路径"].endswith("202606_存款表_1.xlsx")
    assert row["备注"] == "目标同名：将自动后缀"


def test_archive_preview_records_invalid_paths(tmp_path):
    missing = tmp_path / "missing.xlsx"

    stat = run_archive_by_date_preview([missing], tmp_path / "archive", tmp_path / "out", tmp_path / "logs")
    df = _read_plan(stat["output"])

    assert stat["invalid"] == 1
    assert df.iloc[0]["状态"] == "跳过：非文件"


def test_config_init_creates_archive_type_config(tmp_path):
    cfg = tmp_path / "config" / "config.xlsx"

    initialize_or_repair_config(cfg)

    df = pd.read_excel(cfg, sheet_name=SHEET_ARCHIVE_TYPE_CONFIG, dtype=str).fillna("")
    assert list(df.columns) == ARCHIVE_TYPE_CONFIG_COLS
    assert "普惠贷款表" in set(df["匹配关键词"])
    assert "需要归档的后缀" in df.columns
    assert "排除归档的后缀" in df.columns


def test_archive_type_precheck_allows_old_config_without_ext_columns(tmp_path):
    cfg = tmp_path / "config.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_ARCHIVE_TYPE_CONFIG
    for col, name in enumerate(["是否启用", "匹配关键词", "归档文件夹", "备注"], start=1):
        ws.cell(1, col).value = name
    wb.save(cfg)
    wb.close()

    assert validate_sheets_for_feature("ar_type", cfg) == []
    assert validate_sheets_for_feature("ar_type_copy", cfg) == []


def test_archive_ext_filter_applies_to_date_preview(tmp_path):
    cfg = tmp_path / "config.xlsx"
    _make_type_cfg(cfg, include_exts="xlsx;docx", exclude_exts="docx")
    src = tmp_path / "src"
    _touch(src / "20250331_存款表.xlsx")
    _touch(src / "20250331_贷款表.docx")
    _touch(src / "20250331_贷款表.pdf")

    stat = run_archive_by_date_preview([src], tmp_path / "archive", tmp_path / "out", tmp_path / "logs", cfg)
    df = _read_plan(stat["output"])

    by_name = {row["文件名"]: row for _, row in df.iterrows()}
    assert stat["skipped_ext"] == 2
    assert by_name["20250331_存款表.xlsx"]["分类结果"] == "202503"
    assert by_name["20250331_存款表.xlsx"]["状态"] == "待归档"
    assert by_name["20250331_贷款表.docx"]["状态"] == "跳过：后缀被排除"
    assert by_name["20250331_贷款表.pdf"]["状态"] == "跳过：后缀不在归档范围"


def test_archive_ext_filter_applies_to_type_preview(tmp_path):
    cfg = tmp_path / "config.xlsx"
    _make_type_cfg(cfg, exclude_exts="exe")
    src = tmp_path / "src"
    _touch(src / "202503_科技贷款表.xlsx")
    _touch(src / "202503_科技贷款表.exe")

    stat = run_archive_by_type_preview(cfg, [src], tmp_path / "archive", tmp_path / "out", tmp_path / "logs")
    df = _read_plan(stat["output"])

    by_name = {row["文件名"]: row for _, row in df.iterrows()}
    assert stat["skipped_ext"] == 1
    assert by_name["202503_科技贷款表.xlsx"]["分类结果"] == "科技贷款表"
    assert by_name["202503_科技贷款表.exe"]["状态"] == "跳过：后缀被排除"


def test_archive_plan_copy_copies_and_suffixes_existing_target(tmp_path):
    src = tmp_path / "src" / "202606_存款表.xlsx"
    target = tmp_path / "archive" / "202606" / "202606_存款表.xlsx"
    _touch(src)
    _touch(target)
    plan = tmp_path / "plan.xlsx"
    pd.DataFrame([{
        "源文件路径": str(src),
        "文件名": src.name,
        "分类方式": "按日期",
        "分类结果": "202606",
        "目标文件夹": str(target.parent),
        "目标文件路径": str(target),
        "状态": "待归档",
        "备注": "",
    }]).to_excel(plan, sheet_name="归档计划", index=False)

    stat = run_archive_plan_copy(plan, tmp_path / "out", tmp_path / "logs")

    copied = tmp_path / "archive" / "202606" / "202606_存款表_1.xlsx"
    assert stat["copied"] == 1
    assert copied.exists()
    result = pd.read_excel(stat["output"], sheet_name="归档执行结果", dtype=str).fillna("")
    assert result.iloc[0]["执行状态"] == "已复制"
    assert result.iloc[0]["实际目标路径"].endswith("202606_存款表_1.xlsx")


def test_archive_plan_copy_executes_unrecognized_and_skips_invalid_status(tmp_path):
    src = tmp_path / "src" / "未知.xlsx"
    _touch(src)
    target = tmp_path / "archive" / "未识别日期" / "未知.xlsx"
    plan = tmp_path / "plan.xlsx"
    pd.DataFrame([
        {
            "源文件路径": str(src),
            "目标文件路径": str(target),
            "状态": "未识别日期",
        },
        {
            "源文件路径": str(src),
            "目标文件路径": str(tmp_path / "archive" / "skip.xlsx"),
            "状态": "跳过：非文件",
        },
        {
            "源文件路径": str(tmp_path / "missing.xlsx"),
            "目标文件路径": str(tmp_path / "archive" / "missing.xlsx"),
            "状态": "待归档",
        },
    ]).to_excel(plan, sheet_name="归档计划", index=False)

    stat = run_archive_plan_copy(plan, tmp_path / "out", tmp_path / "logs")

    assert stat["copied"] == 1
    assert stat["skip"] == 2
    assert target.exists()
    result = pd.read_excel(stat["output"], sheet_name="归档执行结果", dtype=str).fillna("")
    assert set(result["执行状态"]) == {"已复制", "跳过：计划状态不可执行", "跳过：源文件不存在"}


def test_archive_by_date_copy_copies_to_month_and_unrecognized_dirs(tmp_path):
    src = tmp_path / "src"
    _touch(src / "20260609_存款表.xlsx")
    _touch(src / "无日期_存款表.xlsx")

    stat = run_archive_by_date_copy([src], tmp_path / "archive", tmp_path / "out", tmp_path / "logs")

    assert stat["copied"] == 2
    assert (tmp_path / "archive" / "202606" / "20260609_存款表.xlsx").exists()
    assert (tmp_path / "archive" / "未识别日期" / "无日期_存款表.xlsx").exists()
    result = pd.read_excel(stat["output"], sheet_name="归档执行结果", dtype=str).fillna("")
    assert set(result["执行状态"]) == {"已复制"}


def test_archive_by_date_copy_skips_filtered_extensions(tmp_path):
    cfg = tmp_path / "config.xlsx"
    _make_type_cfg(cfg, include_exts="xlsx", exclude_exts="")
    src = tmp_path / "src"
    _touch(src / "20250331_存款表.xlsx")
    _touch(src / "20250331_安装包.exe")

    stat = run_archive_by_date_copy([src], tmp_path / "archive", tmp_path / "out", tmp_path / "logs", cfg)

    assert stat["copied"] == 1
    assert stat["skip"] == 1
    assert stat["skipped_ext"] == 1
    assert (tmp_path / "archive" / "202503" / "20250331_存款表.xlsx").exists()
    assert not (tmp_path / "archive" / "202503" / "20250331_安装包.exe").exists()


def test_archive_by_type_copy_copies_recursively_and_uses_fallback(tmp_path):
    cfg = tmp_path / "config.xlsx"
    _make_type_cfg(cfg)
    src = tmp_path / "src"
    _touch(src / "nested" / "202606_科技贷款表.xlsx")
    _touch(src / "202606_未知.xlsx")

    stat = run_archive_by_type_copy(cfg, [src], tmp_path / "archive", tmp_path / "out", tmp_path / "logs")

    assert stat["copied"] == 2
    assert stat["rules"] == 3
    assert (tmp_path / "archive" / "科技贷款表" / "202606_科技贷款表.xlsx").exists()
    assert (tmp_path / "archive" / "未识别类型" / "202606_未知.xlsx").exists()
