from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from pytools.core.config_xlsx import SHEET_CONFIG_RENAME
from pytools.core.convert_33_34_35_37 import run_batch_rename_local_map_48


def _make_cfg(path: Path, rows: list[tuple[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_CONFIG_RENAME
    ws["M1"] = "原文件名片段"
    ws["N1"] = "新文件名片段"
    for i, (src, dst) in enumerate(rows, start=2):
        ws.cell(i, 13).value = src
        ws.cell(i, 14).value = dst
    wb.save(path)
    wb.close()


def _touch(path: Path) -> None:
    path.write_text("x", encoding="utf-8")


def test_local_map_renames_any_file_suffix(tmp_path):
    cfg = tmp_path / "config.xlsx"
    logs = tmp_path / "logs"
    xlsx = tmp_path / "测试文件_aa_1.xlsx"
    docx = tmp_path / "测试文件_aa_1.docx"
    pdf = tmp_path / "测试文件_aa_1.pdf"
    _make_cfg(cfg, [("aa", "bb")])
    for p in (xlsx, docx, pdf):
        _touch(p)

    stat = run_batch_rename_local_map_48(cfg, [xlsx, docx, pdf], logs)

    assert stat["ok"] == 3
    assert (tmp_path / "测试文件_bb_1.xlsx").exists()
    assert (tmp_path / "测试文件_bb_1.docx").exists()
    assert (tmp_path / "测试文件_bb_1.pdf").exists()


def test_local_map_replaces_chinese_fragment(tmp_path):
    cfg = tmp_path / "config.xlsx"
    logs = tmp_path / "logs"
    src = tmp_path / "测试文件_aa_1.xlsx"
    _make_cfg(cfg, [("测试", "结果")])
    _touch(src)

    stat = run_batch_rename_local_map_48(cfg, [src], logs)

    assert stat["ok"] == 1
    assert (tmp_path / "结果文件_aa_1.xlsx").exists()


def test_local_map_uses_next_path_when_target_exists(tmp_path):
    cfg = tmp_path / "config.xlsx"
    logs = tmp_path / "logs"
    src = tmp_path / "测试文件_aa_1.xlsx"
    existing = tmp_path / "测试文件_bb_1.xlsx"
    _make_cfg(cfg, [("aa", "bb")])
    _touch(src)
    _touch(existing)

    stat = run_batch_rename_local_map_48(cfg, [src], logs)

    assert stat["ok"] == 1
    assert existing.exists()
    assert (tmp_path / "测试文件_bb_1_1.xlsx").exists()


def test_local_map_skips_when_no_fragment_matched(tmp_path):
    cfg = tmp_path / "config.xlsx"
    logs = tmp_path / "logs"
    src = tmp_path / "测试文件_aa_1.xlsx"
    _make_cfg(cfg, [("cc", "dd")])
    _touch(src)

    stat = run_batch_rename_local_map_48(cfg, [src], logs)

    assert stat["ok"] == 0
    assert stat["skip"] == 1
    assert src.exists()
