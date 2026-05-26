from __future__ import annotations

from datetime import date, datetime

import pandas as pd
from openpyxl import load_workbook

from pytools.core.date_parse import normalize_data_date_value, parse_data_date
from pytools.core.io_excel import append_to_target, write_workbook


def _cell_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    return v


def test_normalize_data_date_value_supports_full_dates_and_month_end():
    assert normalize_data_date_value("20260101") == date(2026, 1, 1)
    assert normalize_data_date_value("2026.01.01") == date(2026, 1, 1)
    assert normalize_data_date_value("2026年1月1日") == date(2026, 1, 1)
    assert normalize_data_date_value("2026.04") == date(2026, 4, 30)
    assert normalize_data_date_value("2026-04") == date(2026, 4, 30)
    assert normalize_data_date_value("202604") == date(2026, 4, 30)
    assert normalize_data_date_value("2024.02") == date(2024, 2, 29)
    assert normalize_data_date_value("2026.02") == date(2026, 2, 28)


def test_parse_data_date_month_only_uses_month_end():
    assert parse_data_date("报表_202604.xlsx")[0] == "2026-04-30"
    assert parse_data_date("报表.xlsx", "2026.04 月报")[0] == "2026-04-30"


def test_write_workbook_formats_data_date_column(tmp_path):
    out = tmp_path / "out.xlsx"
    df = pd.DataFrame([["2026.04", 1], ["20260101", 2]], columns=["数据日期", "数值"])

    write_workbook(out, {"结果": df})

    wb = load_workbook(out, data_only=True)
    ws = wb["结果"]
    assert _cell_date(ws["A2"].value) == date(2026, 4, 30)
    assert ws["A2"].number_format == "yyyy/m/d"
    assert _cell_date(ws["A3"].value) == date(2026, 1, 1)
    assert ws["A3"].number_format == "yyyy/m/d"
    wb.close()


def test_append_to_target_formats_new_data_date_rows(tmp_path):
    target = tmp_path / "target.xlsx"

    append_to_target(
        target,
        "汇总",
        ["数据日期", "行头路径", "数值"],
        [["2026.04", "A", 1]],
        dedup_key_idx=[0, 1],
        required_prefix=["数据日期", "行头路径"],
    )

    wb = load_workbook(target, data_only=True)
    ws = wb["汇总"]
    assert _cell_date(ws["A2"].value) == date(2026, 4, 30)
    assert ws["A2"].number_format == "yyyy/m/d"
    wb.close()
