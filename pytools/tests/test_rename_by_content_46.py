from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from pytools.core.config_xlsx import SHEET_CONFIG_RENAME
from pytools.core.config_init import initialize_or_repair_config
from pytools.core.convert_33_34_35_37 import (
    _load_content_rename_spec,
    run_batch_rename_by_content_46,
    run_batch_unrename_by_content_47,
)


def _make_cfg(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_CONFIG_RENAME
    ws["O1"] = "内容重命名启用"
    ws["P1"] = "内容重命名取值规则"
    ws["Q1"] = "备注"
    ws["O2"] = "Y"
    ws["P2"] = "sheet1@A1;sheet2@A2"
    ws["Q2"] = "测试配置"
    wb.save(path)
    wb.close()


def _make_l2_only_cfg(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_CONFIG_RENAME
    ws["L1"] = "根据文件内容重命名"
    ws["L2"] = "sheet1@A1;sheet2@A2"
    wb.save(path)
    wb.close()


def _make_source(path: Path) -> None:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "sheet1"
    ws1["A1"] = "A"
    ws2 = wb.create_sheet("sheet2")
    ws2["A2"] = "B"
    wb.save(path)
    wb.close()


def test_content_rename_skips_when_all_enabled_config_values_already_in_name(tmp_path):
    cfg = tmp_path / "config.xlsx"
    logs = tmp_path / "logs"
    src = tmp_path / "测试文件.xlsx"
    _make_cfg(cfg)
    _make_source(src)

    stat1 = run_batch_rename_by_content_46(cfg, [src], logs)
    renamed = tmp_path / "A_B_测试文件.xlsx"
    assert stat1["ok"] == 1
    assert renamed.exists()

    stat2 = run_batch_rename_by_content_46(cfg, [renamed], logs)

    assert stat2["ok"] == 0
    assert stat2["skip"] == 1
    assert renamed.exists()
    assert not (tmp_path / "A_B_A_B_测试文件.xlsx").exists()


def test_content_rename_does_not_skip_when_only_some_enabled_config_values_in_name(tmp_path):
    cfg = tmp_path / "config.xlsx"
    logs = tmp_path / "logs"
    src = tmp_path / "A_测试文件.xlsx"
    _make_cfg(cfg)
    _make_source(src)

    stat = run_batch_rename_by_content_46(cfg, [src], logs)

    assert stat["ok"] == 1
    assert (tmp_path / "A_B_A_测试文件.xlsx").exists()


def test_content_unrename_removes_single_prefix(tmp_path):
    cfg = tmp_path / "config.xlsx"
    logs = tmp_path / "logs"
    src = tmp_path / "A_B_测试文件.xlsx"
    _make_cfg(cfg)
    _make_source(src)

    stat = run_batch_unrename_by_content_47(cfg, [src], logs)

    assert stat["ok"] == 1
    assert (tmp_path / "测试文件.xlsx").exists()
    assert not src.exists()


def test_content_unrename_removes_repeated_prefixes(tmp_path):
    cfg = tmp_path / "config.xlsx"
    logs = tmp_path / "logs"
    src = tmp_path / "A_B_A_B_测试文件.xlsx"
    _make_cfg(cfg)
    _make_source(src)

    stat = run_batch_unrename_by_content_47(cfg, [src], logs)

    assert stat["ok"] == 1
    assert (tmp_path / "测试文件.xlsx").exists()
    assert not src.exists()


def test_content_unrename_removes_partial_prefix(tmp_path):
    cfg = tmp_path / "config.xlsx"
    logs = tmp_path / "logs"
    src = tmp_path / "A_测试文件.xlsx"
    _make_cfg(cfg)
    _make_source(src)

    stat = run_batch_unrename_by_content_47(cfg, [src], logs)

    assert stat["ok"] == 1
    assert (tmp_path / "测试文件.xlsx").exists()
    assert not src.exists()


def test_content_unrename_removes_any_enabled_config_items_from_leading_prefix(tmp_path):
    cfg = tmp_path / "config.xlsx"
    logs = tmp_path / "logs"
    src = tmp_path / "B_A_B_测试文件.xlsx"
    _make_cfg(cfg)
    _make_source(src)

    stat = run_batch_unrename_by_content_47(cfg, [src], logs)

    assert stat["ok"] == 1
    assert (tmp_path / "测试文件.xlsx").exists()
    assert not src.exists()


def test_content_unrename_skips_when_no_l2_prefix_item_at_start(tmp_path):
    cfg = tmp_path / "config.xlsx"
    logs = tmp_path / "logs"
    src = tmp_path / "C_D_测试文件.xlsx"
    _make_cfg(cfg)
    _make_source(src)

    stat = run_batch_unrename_by_content_47(cfg, [src], logs)

    assert stat["ok"] == 0
    assert stat["skip"] == 1
    assert src.exists()


def test_content_unrename_uses_next_path_when_target_exists(tmp_path):
    cfg = tmp_path / "config.xlsx"
    logs = tmp_path / "logs"
    src = tmp_path / "A_B_测试文件.xlsx"
    existing = tmp_path / "测试文件.xlsx"
    _make_cfg(cfg)
    _make_source(src)
    _make_source(existing)

    stat = run_batch_unrename_by_content_47(cfg, [src], logs)

    assert stat["ok"] == 1
    assert existing.exists()
    assert (tmp_path / "测试文件_1.xlsx").exists()


def test_content_rename_ignores_legacy_l2_when_no_opq_enabled(tmp_path):
    cfg = tmp_path / "config.xlsx"
    logs = tmp_path / "logs"
    src = tmp_path / "测试文件.xlsx"
    _make_l2_only_cfg(cfg)
    _make_source(src)

    try:
        run_batch_rename_by_content_46(cfg, [src], logs)
    except ValueError as e:
        assert "没有启用的内容重命名配置" in str(e)
    else:
        raise AssertionError("expected missing enabled content rename config")

    assert src.exists()
    assert not (tmp_path / "A_B_测试文件.xlsx").exists()


def test_content_rename_rejects_multiple_enabled_opq_configs(tmp_path):
    cfg = tmp_path / "config.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_CONFIG_RENAME
    ws["O1"] = "内容重命名启用"
    ws["P1"] = "内容重命名取值规则"
    ws["Q1"] = "备注"
    ws["O2"] = "Y"
    ws["P2"] = "sheet1@A1"
    ws["Q2"] = "第一套"
    ws["O3"] = "是"
    ws["P3"] = "sheet2@A2"
    ws["Q3"] = "第二套"
    wb.save(cfg)
    wb.close()

    try:
        _load_content_rename_spec(cfg)
    except ValueError as e:
        msg = str(e)
        assert "只能启用一条" in msg
        assert "第2行" in msg
        assert "第一套" in msg
        assert "第3行" in msg
        assert "第二套" in msg
    else:
        raise AssertionError("expected multiple enabled content rename configs")


def test_config_init_creates_content_rename_opq_config(tmp_path):
    cfg = tmp_path / "config.xlsx"

    initialize_or_repair_config(cfg)

    wb = load_workbook(cfg)
    try:
        ws = wb[SHEET_CONFIG_RENAME]
        assert ws["O1"].value == "内容重命名启用"
        assert ws["P1"].value == "内容重命名取值规则"
        assert ws["Q1"].value == "备注"
        assert ws["O2"].value == "Y"
        assert ws["P2"].value == "sheet1@A1;sheet2@A2"
        assert ws["Q2"].value == "示例：按文件内容单元格生成前缀"
    finally:
        wb.close()


def test_config_repair_adds_opq_config_to_legacy_l2_sheet(tmp_path):
    cfg = tmp_path / "config.xlsx"
    _make_l2_only_cfg(cfg)

    initialize_or_repair_config(cfg)

    wb = load_workbook(cfg)
    try:
        ws = wb[SHEET_CONFIG_RENAME]
        assert ws["L1"].value == "占位L"
        assert ws["O1"].value == "内容重命名启用"
        assert ws["P1"].value == "内容重命名取值规则"
        assert ws["Q1"].value == "备注"
        assert ws["O2"].value == "Y"
        assert ws["P2"].value == "sheet1@A1;sheet2@A2"
        assert ws["Q2"].value == "示例：按文件内容单元格生成前缀"
    finally:
        wb.close()
