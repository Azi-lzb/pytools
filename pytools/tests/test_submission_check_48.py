from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from pytools.core.config_init import initialize_or_repair_config
from pytools.core.config_xlsx import SHEET_SUBMISSION_CHECK_CONFIG
from pytools.core.convert_33_34_35_37 import run_submission_check_48


def _make_cfg(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_SUBMISSION_CHECK_CONFIG
    ws.append(["是否启用", "日期关键词", "县区关键词", "机构关键词", "备用关键词", "命中次数", "命中的文件名", "备注"])
    ws.append(["Y", "202506", "惠城区", "建设银行", "存款表", "", "", "应提交"])
    ws.append(["Y", "202506", "惠阳区", "农业银行", "贷款表", "", "", "应提交"])
    ws.append(["Y", "202506", "", "工商银行", "存款表", "", "", "文件名没有县区也可检查"])
    ws.append(["N", "202506", "博罗县", "中国银行", "存款表", "", "", "禁用行不检查"])
    wb.save(path)
    wb.close()


def test_submission_check_counts_hits_and_outputs_unmatched(tmp_path):
    cfg = tmp_path / "config.xlsx"
    out_dir = tmp_path / "output"
    logs = tmp_path / "logs"
    _make_cfg(cfg)

    files = [
        tmp_path / "202506_惠城区_建设银行_存款表.xlsx",
        tmp_path / "202506_工商银行_存款表.xlsx",
        tmp_path / "随便一个未识别文件.xlsx",
    ]
    for p in files:
        p.write_text("x", encoding="utf-8")

    stat = run_submission_check_48(cfg, files, out_dir, logs)

    assert stat["rules"] == 3
    assert stat["files"] == 3
    assert stat["submitted"] == 2
    assert stat["missing"] == 1
    assert stat["unmatched"] == 1
    assert Path(stat["output"]).exists()

    result = pd.read_excel(stat["output"], sheet_name="检查结果", dtype=object).fillna("")
    assert result.loc[0, "命中次数"] == 1
    assert result.loc[0, "命中的文件名"] == "202506_惠城区_建设银行_存款表.xlsx"
    assert result.loc[1, "命中次数"] == 0
    assert result.loc[2, "命中次数"] == 1
    assert result.loc[2, "命中的文件名"] == "202506_工商银行_存款表.xlsx"

    missing = pd.read_excel(stat["output"], sheet_name="缺失清单", dtype=object).fillna("")
    assert list(missing["机构关键词"]) == ["农业银行"]

    unmatched = pd.read_excel(stat["output"], sheet_name="未识别文件", dtype=object).fillna("")
    assert list(unmatched["文件名"]) == ["随便一个未识别文件.xlsx"]


def test_submission_check_reports_multi_matched_files(tmp_path):
    cfg = tmp_path / "config.xlsx"
    out_dir = tmp_path / "output"
    logs = tmp_path / "logs"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_SUBMISSION_CHECK_CONFIG
    ws.append(["是否启用", "日期关键词", "县区关键词", "机构关键词", "备用关键词", "命中次数", "命中的文件名", "备注"])
    ws.append(["Y", "202506", "", "建设银行", "", "", "", "宽规则"])
    ws.append(["Y", "202506", "惠城区", "建设银行", "", "", "", "窄规则"])
    wb.save(cfg)
    wb.close()
    src = tmp_path / "202506_惠城区_建设银行_存款表.xlsx"
    src.write_text("x", encoding="utf-8")

    stat = run_submission_check_48(cfg, [src], out_dir, logs)

    multi = pd.read_excel(stat["output"], sheet_name="多重命中文件", dtype=object).fillna("")
    assert list(multi["文件名"]) == [src.name]
    assert "第2行" in multi.loc[0, "命中配置"]
    assert "第3行" in multi.loc[0, "命中配置"]


def test_config_init_creates_submission_check_config(tmp_path):
    cfg = tmp_path / "config.xlsx"

    initialize_or_repair_config(cfg)

    wb = load_workbook(cfg)
    try:
        ws = wb[SHEET_SUBMISSION_CHECK_CONFIG]
        headers = [ws.cell(1, c).value for c in range(1, 9)]
        assert headers == ["是否启用", "日期关键词", "县区关键词", "机构关键词", "备用关键词", "命中次数", "命中的文件名", "备注"]
        assert ws["A2"].value == "Y"
        assert ws["B2"].value == "202506"
    finally:
        wb.close()
