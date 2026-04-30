from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .logger import get_logger


BRANCH_SHEET_KEY = "分机构"


def _norm_name(original: str, mapping: dict[str, str], keys_lower: list[tuple[str, str]]) -> str:
    """精确匹配优先；否则不区分大小写做子串匹配。

    keys_lower: [(原始key小写, 原始key), ...] 按插入顺序，用于稳定的子串匹配。
    """
    if original in mapping:
        return mapping[original]
    low = original.lower()
    for k_low, k in keys_lower:
        if k_low and k_low in low:
            return mapping[k]
    return original


def _find_last_row(ws, col: int) -> int:
    for r in range(ws.max_row, 0, -1):
        v = ws.cell(row=r, column=col).value
        if v is not None and str(v).strip() != "":
            return r
    return 0


def _process_sheet(ws, mapping: dict[str, str], keys_lower: list[tuple[str, str]]) -> int:
    """处理单个分机构工作表；返回修改行数。

    机构名可能在 A 列或 B 列：A 非空则以 A 为准；否则 B。
    数据起始行固定为 2。
    """
    last_a = _find_last_row(ws, 1)
    last_b = _find_last_row(ws, 2)
    last_row = max(last_a, last_b)
    if last_row < 2:
        return 0

    def _col_has_data(col: int) -> bool:
        for r in range(2, last_row + 1):
            v = ws.cell(row=r, column=col).value
            if v is not None and str(v).strip() != "":
                return True
        return False

    if _col_has_data(1):
        target_col = 1
    elif _col_has_data(2):
        target_col = 2
    else:
        return 0

    changed = 0
    for r in range(2, last_row + 1):
        v = ws.cell(row=r, column=target_col).value
        if v is None:
            continue
        original = str(v).strip()
        if not original:
            continue
        mapped = _norm_name(original, mapping, keys_lower)
        if mapped != original:
            ws.cell(row=r, column=target_col, value=mapped)
            changed += 1
    return changed


def _process_one_workbook(path: Path, mapping: dict[str, str],
                          keys_lower: list[tuple[str, str]], logger) -> dict:
    stat = {"branch_sheets": 0, "mapped_sheets": 0, "mapped_rows": 0}
    wb = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
    try:
        for ws in wb.worksheets:
            if BRANCH_SHEET_KEY not in ws.title:
                continue
            stat["branch_sheets"] += 1
            n = _process_sheet(ws, mapping, keys_lower)
            if n > 0:
                stat["mapped_sheets"] += 1
                stat["mapped_rows"] += n
                logger.info(f"  sheet '{ws.title}' 修改 {n} 行")
        if stat["mapped_rows"] > 0:
            wb.save(path)
    finally:
        wb.close()
    return stat


def run_normalize_institution(target_wbs: Iterable[Path], mapping: dict[str, str],
                              log_dir: Path) -> dict:
    logger = get_logger("feature_1_2_normalize_institution", log_dir)
    if not mapping:
        logger.warning("机构映射表为空，无法继续")
        return {"files": 0, "branch_sheets": 0, "mapped_sheets": 0,
                "mapped_rows": 0, "failed_files": 0, "mapping_empty": True}

    logger.info(f"机构映射条目数: {len(mapping)}")
    keys_lower = [(k.lower(), k) for k in mapping.keys()]

    total = {"files": 0, "branch_sheets": 0, "mapped_sheets": 0,
             "mapped_rows": 0, "failed_files": 0, "mapping_empty": False}
    for path in target_wbs:
        try:
            logger.info(f"处理文件: {path}")
            stat = _process_one_workbook(path, mapping, keys_lower, logger)
            total["files"] += 1
            for k in ("branch_sheets", "mapped_sheets", "mapped_rows"):
                total[k] += stat[k]
            logger.info(
                f"文件完成: {path.name} 分机构={stat['branch_sheets']} "
                f"命中={stat['mapped_sheets']} 行={stat['mapped_rows']}"
            )
        except Exception as e:
            total["failed_files"] += 1
            logger.error(f"文件失败: {path} - {e}")
    return total
