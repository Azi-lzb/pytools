from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .logger import get_logger


INDICATORS = ("村镇银行本外币", "村镇银行人民币", "村镇银行外汇")
SPLIT_ROW_1_NAME = "博罗长江村镇银行"
SPLIT_ROW_2_NAME = "惠东惠民村镇银行"
VILLAGE_BANK_KEY = "村镇银行"
BRANCH_SHEET_KEY = "分机构"


def _to_num(v) -> float:
    if v is None:
        return 0.0
    try:
        return float(v)
    except (ValueError, TypeError):
        try:
            return float(str(v).strip())
        except (ValueError, TypeError):
            return 0.0


def _yoy(curr: float, prev: float):
    if prev == 0:
        return 0 if curr == 0 else "∞"
    return round((curr / prev - 1) * 100, 2)


def _find_sheet_by_substr(wb, substr: str):
    for name in wb.sheetnames:
        if substr in name:
            return wb[name]
    return None


def compute_village_bank_matrix(current_path: Path, prev_path: Path) -> list[list]:
    """返回 6×8 矩阵：行 1-3 为三类指标 row5 数据（存款 C/D/E+同比，贷款 CI/CJ/CK+同比）；行 4-6 为同三指标 row6 数据。"""
    matrix: list[list] = [[0] * 8 for _ in range(6)]

    wb_cur = load_workbook(current_path, data_only=True, read_only=True)
    wb_prev = load_workbook(prev_path, data_only=True, read_only=True)
    try:
        for idx, indicator in enumerate(INDICATORS, start=1):
            ws_cur = _find_sheet_by_substr(wb_cur, indicator)
            if ws_cur is None:
                continue
            ws_prev = wb_prev[ws_cur.title] if ws_cur.title in wb_prev.sheetnames else None
            if ws_prev is None:
                continue

            for row_offset, base_row in enumerate((5, 6)):
                target_row = idx + row_offset * 3
                # 存款 C/D/E
                curr_dep = _to_num(ws_cur[f"C{base_row}"].value)
                prev_dep = _to_num(ws_prev[f"C{base_row}"].value)
                matrix[target_row - 1][0] = curr_dep
                matrix[target_row - 1][1] = ws_cur[f"D{base_row}"].value
                matrix[target_row - 1][2] = ws_cur[f"E{base_row}"].value
                matrix[target_row - 1][3] = _yoy(curr_dep, prev_dep)
                # 贷款 CI/CJ/CK
                curr_loan = _to_num(ws_cur[f"CI{base_row}"].value)
                prev_loan = _to_num(ws_prev[f"CI{base_row}"].value)
                matrix[target_row - 1][4] = curr_loan
                matrix[target_row - 1][5] = ws_cur[f"CJ{base_row}"].value
                matrix[target_row - 1][6] = ws_cur[f"CK{base_row}"].value
                matrix[target_row - 1][7] = _yoy(curr_loan, prev_loan)
    finally:
        wb_cur.close()
        wb_prev.close()
    return matrix


def _find_village_bank_row(ws) -> int | None:
    for row in range(1, ws.max_row + 1):
        v = ws.cell(row=row, column=1).value
        if v is not None and str(v).strip() == VILLAGE_BANK_KEY:
            return row
    return None


def _detect_sheet_type(name: str) -> str | None:
    if "本外币" in name:
        return "本外币"
    if "人民币" in name:
        return "人民币"
    if "外汇" in name:
        return "外汇"
    return None


def _fill_rows(ws, anchor_row: int, matrix: list[list], sheet_type: str) -> None:
    row_map = {"本外币": (0, 3), "人民币": (1, 4), "外汇": (2, 5)}
    top_idx, bot_idx = row_map[sheet_type]
    for i in range(8):
        ws.cell(row=anchor_row, column=i + 2, value=matrix[top_idx][i])
        ws.cell(row=anchor_row + 1, column=i + 2, value=matrix[bot_idx][i])


def _process_one_workbook(path: Path, matrix: list[list], logger) -> dict:
    stat = {"branch_sheets": 0, "filled_sheets": 0, "skipped_sheets": 0}
    wb = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
    try:
        for ws in wb.worksheets:
            if BRANCH_SHEET_KEY not in ws.title:
                continue
            stat["branch_sheets"] += 1
            sheet_type = _detect_sheet_type(ws.title)
            if sheet_type is None:
                stat["skipped_sheets"] += 1
                logger.info(f"跳过 sheet（未识别类型）: {ws.title}")
                continue
            target_row = _find_village_bank_row(ws)
            if target_row is None:
                stat["skipped_sheets"] += 1
                logger.info(f"跳过 sheet（未找到'{VILLAGE_BANK_KEY}'行）: {ws.title}")
                continue
            ws.delete_rows(target_row, 1)
            ws.insert_rows(target_row, amount=2)
            ws.cell(row=target_row, column=1, value=SPLIT_ROW_1_NAME)
            ws.cell(row=target_row + 1, column=1, value=SPLIT_ROW_2_NAME)
            _fill_rows(ws, target_row, matrix, sheet_type)
            stat["filled_sheets"] += 1
            logger.info(f"已处理 sheet: {ws.title} (行 {target_row}, 类型 {sheet_type})")
        wb.save(path)
    finally:
        wb.close()
    return stat


def run_split_village_bank(current_wb: Path, prev_wb: Path,
                           target_wbs: Iterable[Path], log_dir: Path) -> dict:
    logger = get_logger("feature_1_1_split_village_bank", log_dir)
    logger.info(f"本期文件: {current_wb}")
    logger.info(f"去年同期文件: {prev_wb}")

    matrix = compute_village_bank_matrix(current_wb, prev_wb)
    if all(matrix[0][i] in (0, None) for i in (0, 4)):
        logger.warning("村镇银行计算结果为空（第一行存款与贷款均为 0），中止")
        return {"files": 0, "branch_sheets": 0, "filled_sheets": 0, "skipped_sheets": 0,
                "failed_files": 0, "matrix_empty": True}

    logger.info("计算矩阵：")
    for i, row in enumerate(matrix, 1):
        logger.info(f"  行{i}: {row}")

    total = {"files": 0, "branch_sheets": 0, "filled_sheets": 0, "skipped_sheets": 0,
             "failed_files": 0, "matrix_empty": False}
    for path in target_wbs:
        try:
            stat = _process_one_workbook(path, matrix, logger)
            total["files"] += 1
            for k in ("branch_sheets", "filled_sheets", "skipped_sheets"):
                total[k] += stat[k]
            logger.info(f"文件完成: {path.name} 分机构 {stat['branch_sheets']} 填入 {stat['filled_sheets']}")
        except Exception as e:
            total["failed_files"] += 1
            logger.error(f"文件失败: {path} - {e}")
    return total
