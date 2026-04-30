from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string

from .logger import get_logger


_INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")


def _normalize_range_spec(text: str) -> str:
    if not text:
        return ""
    s = text.replace(" ", "").replace("\t", "")
    for ch in ("，", "、", "；", ";"):
        s = s.replace(ch, ",")
    s = s.replace("：", ":")
    return s


def _token_to_index(token: str, is_row: bool) -> int:
    token = token.strip()
    if not token:
        return 0
    if token.isdigit():
        return int(token)
    if is_row:
        return 0
    try:
        return column_index_from_string(token.upper())
    except Exception:
        return 0


def _parse_indexes(spec: str, max_value: int, is_row: bool) -> list[int]:
    norm = _normalize_range_spec(spec)
    if not norm or max_value <= 0:
        return list(range(1, max_value + 1))
    result: list[int] = []
    for seg in norm.split(","):
        seg = seg.strip()
        if not seg:
            continue
        if ":" in seg:
            parts = seg.split(":", 1)
            start = _token_to_index(parts[0], is_row)
            end = _token_to_index(parts[1], is_row)
            if start <= 0 or end <= 0:
                continue
            if start > end:
                start, end = end, start
            start = max(start, 1)
            end = min(end, max_value)
            result.extend(range(start, end + 1))
        else:
            idx = _token_to_index(seg, is_row)
            if 1 <= idx <= max_value:
                result.append(idx)
    if not result:
        return list(range(1, max_value + 1))
    # 保序去重
    seen: set[int] = set()
    out: list[int] = []
    for v in result:
        if v not in seen:
            seen.add(v)
            out.append(v)
    return out


def _build_exact_name(keywords: list[str]) -> str:
    return "-".join(k for k in keywords if k.strip())


def _sheet_matches_all(name: str, keywords: list[str]) -> bool:
    name_low = name.lower()
    for k in keywords:
        k = k.strip()
        if not k:
            continue
        if k.lower() not in name_low:
            return False
    return True


def _find_matched_sheets(wb, keywords: list[str]) -> list:
    matched = []
    seen: set[str] = set()
    exact = _build_exact_name(keywords)
    if exact and exact in wb.sheetnames:
        matched.append(wb[exact])
        seen.add(exact)
    for ws in wb.worksheets:
        if ws.title in seen:
            continue
        if _sheet_matches_all(ws.title, keywords):
            matched.append(ws)
            seen.add(ws.title)
    return matched


def _sanitize_sheet_name(name: str) -> str:
    cleaned = _INVALID_SHEET_CHARS.sub("_", name)
    return cleaned[:31] if len(cleaned) > 31 else cleaned


def _unique_sheet_name(wb, base: str) -> str:
    base = _sanitize_sheet_name(base)
    if base not in wb.sheetnames:
        return base
    root = base[:28] if len(base) > 28 else base
    for i in range(2, 101):
        cand = f"{root}_{i}"
        if cand not in wb.sheetnames:
            return cand
    return f"{root}_{datetime.now().strftime('%H%M%S')}"


def _read_ws_values(ws) -> list[list]:
    max_r = ws.max_row or 0
    max_c = ws.max_column or 0
    if max_r == 0 or max_c == 0:
        return []
    return [[ws.cell(row=r, column=c).value for c in range(1, max_c + 1)]
            for r in range(1, max_r + 1)]


def _extract_values(data: list[list], rows_spec: str, cols_spec: str,
                    full_extract: bool) -> list[list]:
    if not data:
        return []
    max_r = len(data)
    max_c = max((len(row) for row in data), default=0)
    if max_c == 0:
        return []

    if full_extract or (not rows_spec and not cols_spec):
        rows_idx = list(range(1, max_r + 1))
        cols_idx = list(range(1, max_c + 1))
    else:
        rows_idx = _parse_indexes(rows_spec, max_r, is_row=True)
        cols_idx = _parse_indexes(cols_spec, max_c, is_row=False)

    out: list[list] = []
    for r in rows_idx:
        src_row = data[r - 1] if r - 1 < len(data) else []
        new_row = [src_row[c - 1] if c - 1 < len(src_row) else None for c in cols_idx]
        out.append(new_row)
    return out


def _compact(values: list[list]) -> list[list]:
    """去除全空行与全空列。"""
    def _is_empty(v) -> bool:
        return v is None or (isinstance(v, str) and v.strip() == "")

    rows_kept = [row for row in values if any(not _is_empty(v) for v in row)]
    if not rows_kept:
        return []
    ncol = max(len(r) for r in rows_kept)
    # padding
    rows_kept = [r + [None] * (ncol - len(r)) for r in rows_kept]
    col_has_data = [any(not _is_empty(r[c]) for r in rows_kept) for c in range(ncol)]
    return [[r[c] for c in range(ncol) if col_has_data[c]] for r in rows_kept]


def _write_values(ws_out, values: list[list]) -> None:
    for i, row in enumerate(values, start=1):
        for j, v in enumerate(row, start=1):
            if v is not None:
                ws_out.cell(row=i, column=j, value=v)


def _remove_default_blank(wb) -> None:
    if len(wb.sheetnames) > 1:
        first = wb.worksheets[0]
        if first.max_row == 1 and first.max_column == 1 and first.cell(1, 1).value is None:
            wb.remove(first)


def run_extract_sheets(target_wbs: Iterable[Path], tasks: list[dict],
                       output_dir: Path, log_dir: Path) -> dict:
    logger = get_logger("feature_1_6_extract_sheets", log_dir)
    if not tasks:
        logger.warning("'工作表提取'配置无启用任务")
        return {"files": 0, "extracted_sheets": 0, "output_files": 0,
                "failed_files": 0, "paths": []}

    ts_default = f"提取结果_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    output_dir.mkdir(parents=True, exist_ok=True)

    out_wbs: dict[str, Workbook] = {}
    out_paths: dict[str, Path] = {}
    stat = {"files": 0, "extracted_sheets": 0, "output_files": 0,
            "failed_files": 0, "paths": []}

    for src in target_wbs:
        try:
            logger.info(f"打开源文件: {src}")
            wb_src = load_workbook(src, data_only=True, read_only=False)
        except Exception as e:
            stat["failed_files"] += 1
            logger.error(f"打开失败: {src} - {e}")
            continue

        stat["files"] += 1
        try:
            for task in tasks:
                matched = _find_matched_sheets(wb_src, task["keywords"])
                if not matched:
                    continue

                out_name = task["output_name"] or ts_default
                if out_name not in out_wbs:
                    wb_new = Workbook()
                    out_wbs[out_name] = wb_new
                    out_paths[out_name] = output_dir / f"{out_name}.xlsx"
                wb_out = out_wbs[out_name]

                for ws_src in matched:
                    data = _read_ws_values(ws_src)
                    values = _extract_values(data, task["rows_spec"],
                                             task["cols_spec"], task["full_extract"])
                    values = _compact(values)
                    if not values:
                        logger.info(f"  [skip] {src.name} | {ws_src.title} 无数据")
                        continue

                    sheet_title = _unique_sheet_name(wb_out, ws_src.title)
                    ws_out = wb_out.create_sheet(title=sheet_title)
                    _write_values(ws_out, values)
                    stat["extracted_sheets"] += 1
                    logger.info(
                        f"  提取: {src.name} | {ws_src.title} -> {out_name}.xlsx[{sheet_title}]"
                    )
        finally:
            wb_src.close()

    for name, wb_out in out_wbs.items():
        _remove_default_blank(wb_out)
        path = out_paths[name]
        wb_out.save(path)
        wb_out.close()
        stat["output_files"] += 1
        stat["paths"].append(str(path))
        logger.info(f"已保存: {path}")

    return stat
