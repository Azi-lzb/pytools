from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from collections import defaultdict

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.cell import range_boundaries

from .logger import get_logger


RED_FILL = PatternFill(fill_type="solid", fgColor="FFFF0000")


@dataclass
class DedupTask:
    row_no: int
    enabled: bool
    source_wb: Path
    source_ws: str
    key_cols: list[int]
    target_wb: Path | None = None
    target_ws: str | None = None
    write_mode: str = ""
    task_name: str = ""


def _norm(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _workbook_is_supported(path: Path) -> bool:
    return path.suffix.lower() in (".xlsx", ".xlsm")


def _open_wb(path: Path, read_only: bool = False):
    keep_vba = path.suffix.lower() == ".xlsm"
    return load_workbook(path, read_only=read_only, keep_vba=keep_vba, data_only=False)


def _used_col_count(ws) -> int:
    """快速获取有效列数，避免全表扫描。"""
    try:
        dim = ws.calculate_dimension()
        if dim:
            _, _, max_col, _ = range_boundaries(dim)
            if max_col and max_col > 0:
                return max_col
    except Exception:
        pass
    return ws.max_column or 1


def _get_key_cols_by_comment(ws) -> list[int]:
    cols: list[int] = []
    for c in range(1, ws.max_column + 1):
        cm = ws.cell(1, c).comment
        if cm and "标识列" in _norm(cm.text):
            cols.append(c)
    if cols:
        return cols
    return list(range(1, _used_col_count(ws) + 1))


def _row_key(ws, r: int, key_cols: list[int]) -> tuple[str, ...]:
    return tuple(_norm(ws.cell(r, c).value) for c in key_cols)


def _is_key_blank(key: tuple[str, ...]) -> bool:
    return all(not x for x in key)


def _mark_duplicates(ws, key_cols: list[int], used_col: int | None = None) -> int:
    seen = set()
    if used_col is None:
        used_col = _used_col_count(ws)
    dup_count = 0
    for r in range(2, ws.max_row + 1):
        key = _row_key(ws, r, key_cols)
        if _is_key_blank(key):
            continue
        if key in seen:
            dup_count += 1
            for c in range(1, used_col + 1):
                ws.cell(r, c).fill = RED_FILL
        else:
            seen.add(key)
    return dup_count


def _delete_duplicates(ws, key_cols: list[int]) -> int:
    seen = set()
    to_delete: list[int] = []
    for r in range(2, ws.max_row + 1):
        key = _row_key(ws, r, key_cols)
        if _is_key_blank(key):
            continue
        if key in seen:
            to_delete.append(r)
        else:
            seen.add(key)
    for r in reversed(to_delete):
        ws.delete_rows(r, 1)
    return len(to_delete)


def run_check_by_comment(workbook_path: Path, log_dir: Path) -> dict:
    log = get_logger("3_11_1_check_by_comment", log_dir)
    if not workbook_path.exists():
        raise FileNotFoundError(f"文件不存在: {workbook_path}")
    if not _workbook_is_supported(workbook_path):
        raise ValueError(f"仅支持 xlsx/xlsm: {workbook_path}")

    wb = _open_wb(workbook_path)
    hit_sheets = 0
    marked_rows = 0
    for ws in wb.worksheets:
        a1 = ws["A1"].comment
        if not (a1 and "源数据" in _norm(a1.text)):
            continue
        hit_sheets += 1
        key_cols = _get_key_cols_by_comment(ws)
        marked = _mark_duplicates(ws, key_cols)
        marked_rows += marked
        log.info("sheet=%s key_cols=%s marked=%s", ws.title, ",".join(map(str, key_cols)), marked)

    wb.save(workbook_path)
    wb.close()
    return {"hit_sheets": hit_sheets, "marked_rows": marked_rows}


def run_delete_by_comment(workbook_path: Path, log_dir: Path) -> dict:
    log = get_logger("3_11_2_delete_by_comment", log_dir)
    if not workbook_path.exists():
        raise FileNotFoundError(f"文件不存在: {workbook_path}")
    if not _workbook_is_supported(workbook_path):
        raise ValueError(f"仅支持 xlsx/xlsm: {workbook_path}")

    wb = _open_wb(workbook_path)
    hit_sheets = 0
    deleted_rows = 0
    for ws in wb.worksheets:
        a1 = ws["A1"].comment
        if not (a1 and "源数据" in _norm(a1.text)):
            continue
        hit_sheets += 1
        key_cols = _get_key_cols_by_comment(ws)
        deleted = _delete_duplicates(ws, key_cols)
        deleted_rows += deleted
        log.info("sheet=%s key_cols=%s deleted=%s", ws.title, ",".join(map(str, key_cols)), deleted)

    wb.save(workbook_path)
    wb.close()
    return {"hit_sheets": hit_sheets, "deleted_rows": deleted_rows}


def run_check_by_config(tasks: list[DedupTask], log_dir: Path) -> dict:
    log = get_logger("3_11_3_check_by_config", log_dir)
    groups = defaultdict(list)
    for t in tasks:
        if t.enabled:
            groups[str(t.source_wb)].append(t)

    task_ok = task_skip = marked_total = 0
    for wb_path_text, wb_tasks in groups.items():
        wb_path = Path(wb_path_text)
        if not wb_path.exists() or not _workbook_is_supported(wb_path):
            task_skip += len(wb_tasks)
            log.warning("skip workbook=%s reason=not_found_or_unsupported", wb_path)
            continue
        wb = _open_wb(wb_path)
        changed = False
        for t in wb_tasks:
            if t.source_ws not in wb.sheetnames:
                task_skip += 1
                log.warning("skip row=%s reason=sheet_not_found sheet=%s", t.row_no, t.source_ws)
                continue
            ws = wb[t.source_ws]
            used_col = _used_col_count(ws)
            key_cols = t.key_cols or list(range(1, used_col + 1))
            marked = _mark_duplicates(ws, key_cols, used_col=used_col)
            marked_total += marked
            task_ok += 1
            if marked > 0:
                changed = True
            log.info("ok row=%s sheet=%s marked=%s", t.row_no, t.source_ws, marked)
        if changed:
            wb.save(wb_path)
        wb.close()
    return {"task_ok": task_ok, "task_skip": task_skip, "marked_rows": marked_total}


def run_delete_by_config(tasks: list[DedupTask], log_dir: Path) -> dict:
    log = get_logger("3_11_4_delete_by_config", log_dir)
    groups = defaultdict(list)
    for t in tasks:
        if t.enabled:
            groups[str(t.source_wb)].append(t)

    task_ok = task_skip = deleted_total = 0
    for wb_path_text, wb_tasks in groups.items():
        wb_path = Path(wb_path_text)
        if not wb_path.exists() or not _workbook_is_supported(wb_path):
            task_skip += len(wb_tasks)
            log.warning("skip workbook=%s reason=not_found_or_unsupported", wb_path)
            continue
        wb = _open_wb(wb_path)
        changed = False
        for t in wb_tasks:
            if t.source_ws not in wb.sheetnames:
                task_skip += 1
                log.warning("skip row=%s reason=sheet_not_found sheet=%s", t.row_no, t.source_ws)
                continue
            ws = wb[t.source_ws]
            key_cols = t.key_cols or list(range(1, _used_col_count(ws) + 1))
            deleted = _delete_duplicates(ws, key_cols)
            deleted_total += deleted
            task_ok += 1
            if deleted > 0:
                changed = True
            log.info("ok row=%s sheet=%s deleted=%s", t.row_no, t.source_ws, deleted)
        if changed:
            wb.save(wb_path)
        wb.close()
    return {"task_ok": task_ok, "task_skip": task_skip, "deleted_rows": deleted_total}


def _headers_same(src_ws, tgt_ws) -> bool:
    src_cols = _used_col_count(src_ws)
    tgt_cols = _used_col_count(tgt_ws)
    if tgt_ws.max_row < 1:
        return True
    if tgt_cols < src_cols:
        return False
    for c in range(1, src_cols + 1):
        if _norm(src_ws.cell(1, c).value) != _norm(tgt_ws.cell(1, c).value):
            return False
    return True


def _build_existing_key_set(ws, key_cols: list[int]) -> set[tuple[str, ...]]:
    keys: set[tuple[str, ...]] = set()
    if ws.max_row < 2:
        return keys
    for r in range(2, ws.max_row + 1):
        k = _row_key(ws, r, key_cols)
        if not _is_key_blank(k):
            keys.add(k)
    return keys


def _append_unique_source_to_target(
    src_ws, tgt_ws, key_cols: list[int], existing_keys: set[tuple[str, ...]] | None = None
) -> tuple[int, int, set[tuple[str, ...]]]:
    """增量去重追加：只追加新增，不做删除行。"""
    src_cols = _used_col_count(src_ws)
    src_last_row = src_ws.max_row
    if src_last_row < 1:
        return 0, 0, existing_keys or set()

    # 目标为空时先写入表头
    tgt_is_empty = (tgt_ws.max_row <= 1 and _norm(tgt_ws.cell(1, 1).value) == "")
    if tgt_is_empty:
        for c in range(1, src_cols + 1):
            tgt_ws.cell(1, c).value = src_ws.cell(1, c).value
        out_row = 2
    else:
        out_row = tgt_ws.max_row + 1

    keys = existing_keys if existing_keys is not None else _build_existing_key_set(tgt_ws, key_cols)

    appended = 0
    filtered = 0
    for r in range(2, src_last_row + 1):
        k = tuple(_norm(src_ws.cell(r, c).value) for c in key_cols)
        if not _is_key_blank(k):
            if k in keys:
                filtered += 1
                continue
            keys.add(k)

        for c in range(1, src_cols + 1):
            tgt_ws.cell(out_row, c).value = src_ws.cell(r, c).value
        out_row += 1
        appended += 1

    return appended, filtered, keys


def run_append_by_config(tasks: list[DedupTask], log_dir: Path) -> dict:
    log = get_logger("3_11_5_append_by_config", log_dir)
    task_ok = task_skip = appended_total = deleted_total = 0
    src_wb_cache: dict[str, object] = {}
    tgt_wb_cache: dict[str, object] = {}
    tgt_changed: set[str] = set()
    tgt_key_cache: dict[tuple[str, str, tuple[int, ...]], set[tuple[str, ...]]] = {}

    for t in tasks:
        if not t.enabled:
            continue
        if not t.source_wb.exists() or not _workbook_is_supported(t.source_wb):
            task_skip += 1
            log.warning("skip row=%s reason=source_not_found_or_unsupported", t.row_no)
            continue
        if t.target_wb is None or t.target_ws is None:
            task_skip += 1
            log.warning("skip row=%s reason=target_missing", t.row_no)
            continue

        src_wb_key = str(t.source_wb)
        src_wb = src_wb_cache.get(src_wb_key)
        if src_wb is None:
            src_wb = _open_wb(t.source_wb)
            src_wb_cache[src_wb_key] = src_wb
        if t.source_ws not in src_wb.sheetnames:
            task_skip += 1
            log.warning("skip row=%s reason=source_sheet_not_found sheet=%s", t.row_no, t.source_ws)
            continue
        src_ws = src_wb[t.source_ws]

        tgt_wb_key = str(t.target_wb)
        tgt_wb = tgt_wb_cache.get(tgt_wb_key)
        if tgt_wb is None and t.target_wb.exists():
            if not _workbook_is_supported(t.target_wb):
                task_skip += 1
                log.warning("skip row=%s reason=target_unsupported", t.row_no)
                continue
            tgt_wb = _open_wb(t.target_wb)
            tgt_wb_cache[tgt_wb_key] = tgt_wb
        elif tgt_wb is None:
            from openpyxl import Workbook
            tgt_wb = Workbook()
            default_ws = tgt_wb.active
            if default_ws is not None and default_ws.title == "Sheet":
                tgt_wb.remove(default_ws)
            tgt_wb_cache[tgt_wb_key] = tgt_wb

        if t.target_ws in tgt_wb.sheetnames:
            tgt_ws = tgt_wb[t.target_ws]
        else:
            tgt_ws = tgt_wb.create_sheet(title=t.target_ws)

        if tgt_ws.max_row >= 1 and _norm(tgt_ws.cell(1, 1).value) != "":
            if not _headers_same(src_ws, tgt_ws):
                task_skip += 1
                log.warning("skip row=%s reason=header_mismatch", t.row_no)
                continue

        key_cols = t.key_cols or list(range(1, _used_col_count(src_ws) + 1))
        key_cache_key = (tgt_wb_key, t.target_ws, tuple(key_cols))
        existing_keys = tgt_key_cache.get(key_cache_key)
        appended, deleted, latest_keys = _append_unique_source_to_target(src_ws, tgt_ws, key_cols, existing_keys)
        tgt_key_cache[key_cache_key] = latest_keys
        if appended > 0:
            tgt_changed.add(tgt_wb_key)

        appended_total += appended
        deleted_total += deleted
        task_ok += 1
        log.info("ok row=%s appended=%s dedup_deleted=%s", t.row_no, appended, deleted)

    # 统一落盘：仅保存发生过新增写入的目标工作簿
    for wb_path_text, wb in tgt_wb_cache.items():
        wb_path = Path(wb_path_text)
        if wb_path_text in tgt_changed:
            wb_path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(wb_path)
        wb.close()

    for wb in src_wb_cache.values():
        wb.close()

    return {
        "task_ok": task_ok,
        "task_skip": task_skip,
        "appended_rows": appended_total,
        "deleted_rows": deleted_total,
    }


def run_precheck_by_config(tasks: list[DedupTask], log_dir: Path) -> dict:
    log = get_logger("3_11_6_precheck_by_config", log_dir)
    ok = warn = fail = 0
    for t in tasks:
        if not t.enabled:
            continue
        if not t.source_wb:
            fail += 1
            log.error("row=%s fail=source_wb_empty", t.row_no)
            continue
        if not t.source_ws:
            fail += 1
            log.error("row=%s fail=source_ws_empty", t.row_no)
            continue
        if not t.source_wb.exists():
            fail += 1
            log.error("row=%s fail=source_wb_not_found", t.row_no)
            continue
        if not _workbook_is_supported(t.source_wb):
            fail += 1
            log.error("row=%s fail=source_wb_unsupported", t.row_no)
            continue
        try:
            wb = _open_wb(t.source_wb, read_only=True)
            exists = t.source_ws in wb.sheetnames
            wb.close()
            if not exists:
                fail += 1
                log.error("row=%s fail=source_ws_not_found", t.row_no)
                continue
        except Exception as e:
            fail += 1
            log.error("row=%s fail=open_source_error err=%s", t.row_no, e)
            continue

        if not t.key_cols:
            warn += 1
            log.warning("row=%s warn=key_cols_empty_use_all_columns", t.row_no)
        ok += 1

    return {"ok": ok, "warn": warn, "fail": fail}
