from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from .logger import get_logger


INST_SHEET_KEYS = ("惠州市", "涉农贷款分机构")
AREA_SHEET_KEYS = ("本外币", "涉农贷款分地区")
HEADER_TEXT = "涉农贷款比上月"
CITY_NAME = "惠州市"

HEADER_SCAN_ROWS = 10
NAME_SCAN_COLS = 3
INST_DATA_START_ROW = 3
INST_NAME_COL = 2
INST_VALUE_COL = 3
INST_DIFF_COL = 4
TOLERANCE = 0.01

_WS_CLEAN = re.compile(r"[\r\n\t　 ]")


def _normalize(v) -> str:
    if v is None:
        return ""
    s = _WS_CLEAN.sub(" ", str(v)).strip()
    while "  " in s:
        s = s.replace("  ", " ")
    return s


def _to_num(v) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "-"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def _last_row(ws) -> int:
    for r in range(ws.max_row or 0, 0, -1):
        for c in range(1, (ws.max_column or 0) + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None and str(v).strip() != "":
                return r
    return 0


def _last_col(ws) -> int:
    for c in range(ws.max_column or 0, 0, -1):
        for r in range(1, (ws.max_row or 0) + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None and str(v).strip() != "":
                return c
    return 0


def _find_sheet(wb, keys: tuple[str, ...]):
    hits = [ws for ws in wb.worksheets if all(k in ws.title for k in keys)]
    if len(hits) > 1:
        raise ValueError(
            f"工作簿“{Path(wb.path).name if getattr(wb,'path',None) else ''}”"
            f"匹配到多个目标工作表（关键词 {keys}），请保留唯一目标表。"
        )
    return hits[0] if hits else None


def _build_inst_map(ws) -> tuple[dict[str, float], dict[str, int], list[str]]:
    values: dict[str, float] = {}
    rows: dict[str, int] = {}
    dups: list[str] = []
    last_r = _last_row(ws)
    for r in range(INST_DATA_START_ROW, last_r + 1):
        name = _normalize(ws.cell(row=r, column=INST_NAME_COL).value)
        if not name:
            continue
        if name in values:
            dups.append(name)
            continue
        values[name] = _to_num(ws.cell(row=r, column=INST_VALUE_COL).value)
        rows[name] = r
    return values, rows, dups


def _find_header_col(ws, header: str) -> int:
    last_col = _last_col(ws)
    max_r = min(_last_row(ws), HEADER_SCAN_ROWS)
    for r in range(1, max_r + 1):
        for c in range(1, last_col + 1):
            if _normalize(ws.cell(row=r, column=c).value) == header:
                return c
    return 0


def _find_city_row(ws, city: str) -> int:
    last_r = _last_row(ws)
    for r in range(1, last_r + 1):
        for c in range(1, NAME_SCAN_COLS + 1):
            if _normalize(ws.cell(row=r, column=c).value) == city:
                return r
    return 0


def run_adjust_rural_loan(current_path: Path, previous_path: Path, log_dir: Path) -> dict:
    logger = get_logger("feature_1_8_adjust_rural_loan", log_dir)
    logger.info(f"本期: {current_path}")
    logger.info(f"上期: {previous_path}")

    result = {
        "matched": 0, "fill_total": 0.0, "area_total": 0.0,
        "unmatched_current": [], "previous_only": [],
        "consistent": False, "error": None,
    }

    keep_vba = current_path.suffix.lower() == ".xlsm"
    wb_cur = load_workbook(current_path, keep_vba=keep_vba)
    wb_prev = load_workbook(previous_path, data_only=True, read_only=True)
    try:
        ws_cur_inst = _find_sheet(wb_cur, INST_SHEET_KEYS)
        if ws_cur_inst is None:
            raise ValueError("本期文件未找到包含“惠州市”且包含“涉农贷款分机构”的工作表。")
        ws_prev_inst = _find_sheet(wb_prev, INST_SHEET_KEYS)
        if ws_prev_inst is None:
            raise ValueError("上期文件未找到包含“惠州市”且包含“涉农贷款分机构”的工作表。")
        ws_cur_area = _find_sheet(wb_cur, AREA_SHEET_KEYS)
        if ws_cur_area is None:
            raise ValueError("本期文件未找到包含“本外币”且包含“涉农贷款分地区”的工作表。")

        cur_values, cur_rows, cur_dups = _build_inst_map(ws_cur_inst)
        if cur_dups:
            raise ValueError("本期分机构表存在重复机构名称：" + "、".join(cur_dups))
        prev_values, _, prev_dups = _build_inst_map(ws_prev_inst)
        if prev_dups:
            raise ValueError("上期分机构表存在重复机构名称：" + "、".join(prev_dups))

        fill_total = 0.0
        unmatched: list[str] = []
        for name, r in cur_rows.items():
            if name in prev_values:
                diff = cur_values[name] - prev_values[name]
                ws_cur_inst.cell(row=r, column=INST_DIFF_COL, value=diff)
                fill_total += diff
                result["matched"] += 1
            else:
                ws_cur_inst.cell(row=r, column=INST_DIFF_COL, value=None)
                unmatched.append(name)

        previous_only = [n for n in prev_values if n not in cur_values]

        header_col = _find_header_col(ws_cur_area, HEADER_TEXT)
        if header_col == 0:
            raise ValueError('未在“本外币涉农贷款分地区”表中找到表头“涉农贷款比上月”。')
        city_row = _find_city_row(ws_cur_area, CITY_NAME)
        if city_row == 0:
            raise ValueError('未在“本外币涉农贷款分地区”表中找到“惠州市”所在行。')
        area_total = _to_num(ws_cur_area.cell(row=city_row, column=header_col).value)

        wb_cur.save(current_path)

        result["fill_total"] = fill_total
        result["area_total"] = area_total
        result["unmatched_current"] = unmatched
        result["previous_only"] = previous_only
        result["consistent"] = abs(fill_total - area_total) <= TOLERANCE

        logger.info(f"匹配机构数={result['matched']}")
        logger.info(f"分机构合计={fill_total:.4f}  分地区总数={area_total:.4f}  "
                    f"{'一致' if result['consistent'] else '不一致'}")
        if unmatched:
            logger.warning(f"本期未匹配到上期: {'、'.join(unmatched)}")
        if previous_only:
            logger.warning(f"上期独有: {'、'.join(previous_only)}")
    except Exception as e:
        result["error"] = str(e)
        logger.error(f"失败: {e}")
    finally:
        wb_prev.close()
        wb_cur.close()
    return result
