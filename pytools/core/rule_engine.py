from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator
from functools import lru_cache
import re
import pandas as pd
from openpyxl import load_workbook

from .config_xlsx import TimelineRule, PathMapRule
from .io_excel import list_sheet_names, read_sheet_2d
from .path_map import standardize, _match_kw
from .date_parse import parse_data_date


@dataclass
class ExtractedCell:
    source_file: str
    sheet_name: str
    rule_name: str
    data_date: str
    row_path: str
    col_path: str
    value: object
    cell_addr: str = ""


def _col_letter(c: int) -> str:
    s = ""
    n = c
    while n > 0:
        s = chr(65 + (n - 1) % 26) + s
        n = (n - 1) // 26
    return s


def _cell_addr(r: int, c: int) -> str:
    return f"{_col_letter(c)}{r}"


def _fallback_col_path(c0: int) -> str:
    """0-based 列号 -> 兜底列头名（列_B / 列_AA）"""
    return f"列_{_col_letter(c0 + 1)}"


def match_all_keywords(source: str, keywords: str) -> bool:
    """工作簿/工作表关键字按 ; / ， / , 分割；全部包含才算匹配（VBA MatchAllKeywords 语义）。"""
    if not keywords:
        return True
    src = (source or "")
    src_lower = src.lower()
    raw = keywords.replace("，", ",").replace("；", ";").replace(",", ";")
    for kw in raw.split(";"):
        kw = kw.strip()
        if kw and kw.lower() not in src_lower:
            return False
    return True


# 旧名兼容
_kw_match = match_all_keywords


def pick_rules_for_workbook(rules: list[TimelineRule], wb_name: str, mode: str) -> list[TimelineRule]:
    """按工作簿命中策略筛规则。mode=all_match/first_match。"""
    matched = [r for r in rules if _kw_match(wb_name, r.wb_keyword)]
    if mode == "first_match":
        return matched[:1]
    return matched


def parse_col_spec(v) -> int:
    """B → 2; '3' → 3; 3 → 3; '' → 0。"""
    if v is None:
        return 0
    if isinstance(v, float) and pd.isna(v):
        return 0
    s = str(v).strip()
    if not s:
        return 0
    if s.replace(".", "", 1).isdigit():
        try:
            return int(float(s))
        except ValueError:
            return 0
    n = 0
    for ch in s.upper():
        if not ("A" <= ch <= "Z"):
            return 0
        n = n * 26 + (ord(ch) - 64)
    return n


def parse_col_specs(v) -> list[int]:
    """解析多列配置：支持 `A,C` / `1,3` / `A;C` / `A，C`；返回正整数列号列表。"""
    if v is None:
        return []
    if isinstance(v, float) and pd.isna(v):
        return []
    s = str(v).strip()
    if not s:
        return []
    parts = (
        s.replace("，", ",")
        .replace("；", ",")
        .replace(";", ",")
        .split(",")
    )
    out: list[int] = []
    seen: set[int] = set()
    for p in parts:
        n = parse_col_spec(p)
        if n > 0 and n not in seen:
            out.append(n)
            seen.add(n)
    return out


def _skip_match(skip_kws: list[str], text: str) -> bool:
    if not skip_kws or not text:
        return False
    return any(k in text for k in skip_kws)


def _required_match(required: list[str], values: list[str]) -> bool:
    if not required:
        return True
    text = "|".join(str(v) for v in values if v is not None)
    return all(rk in text for rk in required)


def _val_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and pd.isna(v):
        return ""
    s = str(v).strip()
    return s


@lru_cache(maxsize=32)
def _load_merge_ranges_file(path_text: str) -> dict[str, tuple[tuple[int, int, int, int], ...]]:
    """一次性读整本工作簿，收集所有 sheet 的合并区域。
    返回 sheet_name -> tuple(min_r0, max_r0, min_c0, max_c0)，0-based。"""
    p = Path(path_text)
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        keep_vba = ext == ".xlsm"
        wb = load_workbook(p, read_only=False, keep_vba=keep_vba, data_only=False)
        try:
            out: dict[str, tuple[tuple[int, int, int, int], ...]] = {}
            for ws in wb.worksheets:
                rngs = tuple(
                    (rg.min_row - 1, rg.max_row - 1, rg.min_col - 1, rg.max_col - 1)
                    for rg in ws.merged_cells.ranges
                )
                out[ws.title] = rngs
            return out
        finally:
            wb.close()
    return _load_merge_ranges_file_by_com(path_text)


def _load_merge_ranges_file_by_com(path_text: str) -> dict[str, tuple[tuple[int, int, int, int], ...]]:
    """openpyxl 不支持的格式（如 .xls）时，回退用 COM 抽取合并区域。"""
    try:
        import win32com.client  # type: ignore
    except Exception:
        return {}

    app = None
    wb = None
    try:
        app = win32com.client.DispatchEx("Excel.Application")
        app.Visible = False
        app.DisplayAlerts = False
        wb = app.Workbooks.Open(path_text, ReadOnly=True, UpdateLinks=0, AddToMru=False)

        out: dict[str, tuple[tuple[int, int, int, int], ...]] = {}
        for ws in wb.Worksheets:
            sheet_name = str(ws.Name)
            ranges: list[tuple[int, int, int, int]] = []
            used = ws.UsedRange
            try:
                merge_areas = used.MergeAreas
                for ma in merge_areas:
                    ranges.append((
                        int(ma.Row) - 1,
                        int(ma.Row + ma.Rows.Count - 1) - 1,
                        int(ma.Column) - 1,
                        int(ma.Column + ma.Columns.Count - 1) - 1,
                    ))
            except Exception:
                # 个别版本/引擎下 MergeAreas 不可用，回退逐单元格扫描
                for row in used.Rows:
                    for cell in row.Cells:
                        try:
                            if bool(cell.MergeCells):
                                ma = cell.MergeArea
                                ranges.append((
                                    int(ma.Row) - 1,
                                    int(ma.Row + ma.Rows.Count - 1) - 1,
                                    int(ma.Column) - 1,
                                    int(ma.Column + ma.Columns.Count - 1) - 1,
                                ))
                        except Exception:
                            continue
            # 去重，避免同一合并区重复加入
            out[sheet_name] = tuple(sorted(set(ranges)))
        return out
    except Exception:
        return {}
    finally:
        try:
            if wb is not None:
                wb.Close(SaveChanges=False)
        except Exception:
            pass
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass


def _load_merge_ranges(path_text: str, sheet_name: str) -> tuple[tuple[int, int, int, int], ...]:
    return _load_merge_ranges_file(path_text).get(sheet_name, tuple())


def _resolve_merged_top_left(df, r: int, c: int, merge_ranges: tuple[tuple[int, int, int, int], ...]) -> str:
    """若 (r,c) 落在合并区域，返回该区域左上角文本；否则返回空。"""
    for min_r, max_r, min_c, max_c in merge_ranges:
        if min_r <= r <= max_r and min_c <= c <= max_c:
            return _val_str(df.iat[min_r, min_c])
    return ""


def _header_cell_text(df, r: int, c: int, merge_ranges: tuple[tuple[int, int, int, int], ...]) -> str:
    """表头取值：先取自身；若空且为合并区域成员则取左上角。"""
    if r < 0 or c < 0 or r >= df.shape[0] or c >= df.shape[1]:
        return ""
    v = _val_str(df.iat[r, c])
    if v:
        return v
    return _resolve_merged_top_left(df, r, c, merge_ranges)


def _header_cell_text_with_left_fill(df, r: int, c: int,
                                     merge_ranges: tuple[tuple[int, int, int, int], ...]) -> str:
    """列表头兜底：若当前为空且合并信息缺失，向左回溯最近非空值（模拟横向合并展开）。"""
    v = _header_cell_text(df, r, c, merge_ranges)
    if v:
        return v
    k = c - 1
    while k >= 0:
        lv = _header_cell_text(df, r, k, merge_ranges)
        if lv:
            return lv
        k -= 1
    return ""


_TAIL_DIGITS_RE = re.compile(r"^(.*?)(\d+)$")


def _normalize_header_noise(curr: str, prev: str) -> str:
    """清洗 .xls 常见列头噪声：若当前仅比前一列多尾部数字（如 境内存款1/11），归并为前一列。"""
    if not curr or not prev:
        return curr
    m = _TAIL_DIGITS_RE.match(curr)
    if not m:
        return curr
    base = (m.group(1) or "").strip()
    if base and base == prev.strip():
        return prev
    return curr


def _strip_tail_digits(text: str) -> str:
    if not text:
        return text
    m = _TAIL_DIGITS_RE.match(text)
    if not m:
        return text
    base = (m.group(1) or "").strip()
    return base or text


def _normalize_header_row_noise(texts: list[str]) -> list[str]:
    """整行修正 .xls 表头污染。

    典型脏值：
    - 境内存款 / 境内存款1 / 境内存款11
    - 活期存款2 / 活期存款21 / 活期存款211

    若一段连续列满足“后一列 = 前一列 + 纯数字后缀”，则整段统一归并为去尾数后的 base。
    """
    if not texts:
        return texts
    out = list(texts)
    n = len(out)
    i = 0
    while i < n:
        curr = out[i]
        if not curr:
            i += 1
            continue
        base = _strip_tail_digits(curr)
        j = i + 1
        prev = curr
        chain_ok = False
        while j < n and out[j]:
            nxt = out[j]
            if _strip_tail_digits(nxt) != base:
                break
            if nxt.startswith(prev) and nxt != prev and nxt[len(prev):].isdigit():
                chain_ok = True
                prev = nxt
                j += 1
                continue
            break
        if chain_ok:
            for k in range(i, j):
                out[k] = base
            i = j
        else:
            i += 1
    return out


def _row_header_text(df, r: int, c: int, merge_ranges: tuple[tuple[int, int, int, int], ...]) -> str:
    """行头取值：先取自身；若空且为合并区域成员则取左上角。"""
    if r < 0 or c < 0 or r >= df.shape[0] or c >= df.shape[1]:
        return ""
    v = _val_str(df.iat[r, c])
    if v:
        return v
    return _resolve_merged_top_left(df, r, c, merge_ranges)


def iter_matching_sheets(source_path: Path, rules: list[TimelineRule]
                         ) -> Iterator[tuple[TimelineRule, str]]:
    sheet_names = list_sheet_names(source_path)
    for rule in rules:
        for sn in sheet_names:
            if not _kw_match(sn, rule.sheet_keyword):
                continue
            yield rule, sn


def extract_cells(source_path: Path, rule: TimelineRule, sheet_name: str,
                  path_maps: list[PathMapRule],
                  row_suffix_enabled: bool) -> list[ExtractedCell]:
    """按规则定位区域并展开成 ExtractedCell 列表（pandas 版：内部用 pd.read_excel）。"""
    df = read_sheet_2d(source_path, sheet_name)
    merge_ranges = _load_merge_ranges(str(source_path), sheet_name)
    return extract_cells_from_arrays(df, merge_ranges, source_path, rule,
                                     sheet_name, path_maps, row_suffix_enabled)


def extract_cells_from_arrays(df, merge_ranges, source_path: Path, rule: TimelineRule,
                              sheet_name: str, path_maps: list[PathMapRule],
                              row_suffix_enabled: bool) -> list[ExtractedCell]:
    """与 extract_cells 同口径，但数据源为已加载的 df（pandas.DataFrame 或 com_engine.ArrayLike）
    + merge_ranges 合并区域元组。供 COM 引擎复用同一套计算逻辑。"""
    rows, cols = df.shape
    if rows == 0 or cols == 0:
        return []

    a1 = _val_str(df.iat[0, 0]) if rows > 0 and cols > 0 else ""
    data_date, _src = parse_data_date(source_path.name, a1)
    # 日期缺失时不跳过整表，沿用空日期继续提取，避免“规则全不命中”的假象
    if not data_date:
        data_date = ""

    row_header_specified = getattr(rule, "row_header_col_specified", True)
    rh_cols_cfg = getattr(rule, "row_header_cols", None)
    if row_header_specified:
        if rh_cols_cfg:
            rh_cols = [c - 1 for c in rh_cols_cfg if c > 0]
        else:
            rh_cols = [rule.row_header_col - 1]
    else:
        rh_cols = []
    ch_rows = [r - 1 for r in rule.col_header_rows]
    if row_header_specified:
        if not rh_cols:
            return []
        if any(c < 0 or c >= cols for c in rh_cols):
            return []
    if any(r < 0 or r >= rows for r in ch_rows):
        return []

    r_start = rule.data_row_start - 1 if rule.data_row_start else max(ch_rows) + 1
    r_end = (rule.data_row_end - 1) if rule.data_row_end else rows - 1
    c_start = rule.data_col_start - 1 if rule.data_col_start else ((max(rh_cols) + 1) if row_header_specified else 0)
    c_end = (rule.data_col_end - 1) if rule.data_col_end else cols - 1

    if r_start > r_end or c_start > c_end:
        return []

    # 预计算列表头文本，并先按整行做一次脏值归一化。
    header_row_texts: dict[int, list[str]] = {}
    for r in ch_rows:
        row_texts = [
            _header_cell_text_with_left_fill(df, r, c, merge_ranges)
            for c in range(c_start, c_end + 1)
        ]
        header_row_texts[r] = _normalize_header_row_noise(row_texts)

    # 必含行头/列头 校验
    rh_values = []
    if row_header_specified:
        for r in range(r_start, r_end + 1):
            parts = [_row_header_text(df, r, c, merge_ranges) for c in rh_cols]
            parts = [p for p in parts if p]
            rh_values.append("_".join(parts))
    ch_values_flat = []
    for r in ch_rows:
        ch_values_flat.extend(header_row_texts.get(r, []))
    if row_header_specified and not _required_match(rule.required_row_headers, rh_values):
        return []
    if not _required_match(rule.required_col_headers, ch_values_flat):
        return []

    # 生成列头路径（按列）
    col_path_raw: dict[int, str] = {}
    for c in range(c_start, c_end + 1):
        parts: list[str] = []
        c0 = c - c_start
        for r in ch_rows:
            row_texts = header_row_texts.get(r, [])
            p = row_texts[c0] if 0 <= c0 < len(row_texts) else ""
            if p:
                parts.append(p)
        col_path_raw[c] = "_".join(parts)

    # 列头同名编号（始终启用）
    col_path_final: dict[int, str] = {}
    seen_col: dict[str, int] = {}
    for c, p in col_path_raw.items():
        if not p:
            col_path_final[c] = _fallback_col_path(c)
            continue
        seen_col[p] = seen_col.get(p, 0) + 1
        col_path_final[c] = f"{p}_{seen_col[p]}" if seen_col[p] > 1 else p
        # 但若该路径在整个 sheet 只出现一次，保持原样
    # 重新核对——若某 raw path 全程只出现一次，去掉 _1
    raw_count: dict[str, int] = {}
    for p in col_path_raw.values():
        if p:
            raw_count[p] = raw_count.get(p, 0) + 1
    for c, p in list(col_path_final.items()):
        raw = col_path_raw[c]
        if raw and raw_count[raw] == 1:
            col_path_final[c] = raw

    # 行头同名处理
    row_path_final: dict[int, str] = {}
    seen_row: dict[str, int] = {}
    raw_row_count: dict[str, int] = {}
    for r in range(r_start, r_end + 1):
        if row_header_specified:
            parts = [_row_header_text(df, r, c, merge_ranges) for c in rh_cols]
            parts = [p for p in parts if p]
            p = "_".join(parts)
        else:
            p = f"ROW#{r + 1}"
        if p:
            raw_row_count[p] = raw_row_count.get(p, 0) + 1
    for r in range(r_start, r_end + 1):
        if row_header_specified:
            parts = [_row_header_text(df, r, c, merge_ranges) for c in rh_cols]
            parts = [p for p in parts if p]
            raw = "_".join(parts)
        else:
            raw = f"ROW#{r + 1}"
        if not raw:
            row_path_final[r] = ""
            continue
        if _skip_match(rule.skip_keywords, raw):
            row_path_final[r] = ""  # 标记跳过
            continue
        if row_suffix_enabled and raw_row_count[raw] > 1:
            seen_row[raw] = seen_row.get(raw, 0) + 1
            row_path_final[r] = f"{raw}_{seen_row[raw]}"
        else:
            row_path_final[r] = raw

    wb_name = source_path.name

    # 预过滤 path_maps：按 (rule, wb, sheet, target) 一次过滤，替换时只做字符串比较，
    # 避免在内层循环对 100+ 条规则逐条做 kw_match 子串匹配。
    def _filter_maps(target: str) -> list[PathMapRule]:
        out_maps: list[PathMapRule] = []
        for m in path_maps:
            if m.applicable_rule and m.applicable_rule != rule.name:
                continue
            if m.target not in (target, "两者", ""):
                continue
            if not _match_kw(m.wb_keyword, wb_name):
                continue
            if not _match_kw(m.sheet_keyword, sheet_name):
                continue
            out_maps.append(m)
        return out_maps

    def _apply_maps(path: str, maps: list[PathMapRule]) -> str:
        if not path:
            return path
        for m in maps:
            if m.match_mode == "包含":
                if m.original and m.original in path:
                    return path.replace(m.original, m.standard)
            else:
                if path == m.original:
                    return m.standard
        return path

    row_maps = _filter_maps("行头")
    col_maps = _filter_maps("列头")

    # 列头只依赖列，先预计算，避免放入 row×col 内层循环反复算
    col_std_cache: dict[int, str] = {
        c: _apply_maps(col_path_final[c], col_maps) for c in range(c_start, c_end + 1)
    }

    # 用 numpy 数组定位比 df.iat 快很多
    values = df.values

    out: list[ExtractedCell] = []
    for r in range(r_start, r_end + 1):
        rp = row_path_final[r]
        if not rp:
            continue
        rp_std = _apply_maps(rp, row_maps)
        row_values = values[r]
        for c in range(c_start, c_end + 1):
            cp_std = col_std_cache[c]
            if not cp_std:
                continue
            v = row_values[c]
            if v is None or (isinstance(v, float) and pd.isna(v)):
                continue
            out.append(ExtractedCell(
                source_file=wb_name,
                sheet_name=sheet_name,
                rule_name=rule.name,
                data_date=data_date,
                row_path=rp_std,
                col_path=cp_std,
                value=v,
                cell_addr=_cell_addr(r + 1, c + 1),
            ))
    return out
