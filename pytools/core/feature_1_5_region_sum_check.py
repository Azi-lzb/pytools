from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .logger import get_logger


REGION_SHEET_KEY = "分地区"
TOTAL_REGION = "惠州市"
SUB_REGIONS = ("惠阳区", "博罗县", "惠东县", "龙门县", "惠州市本部", "大亚湾")
ABS_TOLERANCE = 0.01
REL_TOLERANCE = 0.001

RESULT_HEADERS = ["序号", "文件名", "工作表名", "核对状态", "错误数量", "错误信息"]
HEADER_FILL = PatternFill("solid", fgColor="C8DCF0")
ERROR_FILL = PatternFill("solid", fgColor="FFC8C8")
ERROR_FONT = Font(color="C00000")
OK_FILL = PatternFill("solid", fgColor="C8FFC8")
OK_FONT = Font(color="006400")


def _to_num_or_none(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s == "" or s == "-":
        return 0.0
    try:
        return float(s)
    except ValueError:
        return None


def _read_all(ws) -> list[list]:
    """一次性流式读取所有单元格值，避免 read_only 下 ws.cell 的 O(N²) 成本。"""
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _check_one_sheet(ws) -> tuple[int, str]:
    """返回 (错误数, 错误详情)。"""
    data = _read_all(ws)
    if not data or not data[0]:
        return 1, "工作表为空或没有数据"

    names_wanted = {TOTAL_REGION, *SUB_REGIONS}
    rows: dict[str, int] = {}
    for r, row in enumerate(data, start=1):
        if len(row) < 2:
            continue
        v = row[1]
        if v is None:
            continue
        name = str(v).strip()
        if name in names_wanted and name not in rows:
            rows[name] = r

    if TOTAL_REGION not in rows:
        return 1, f"未找到'{TOTAL_REGION}'行"

    total_row_idx = rows[TOTAL_REGION] - 1
    total_row_data = data[total_row_idx]
    last_col = len(total_row_data)
    start_col = 3
    if last_col < start_col:
        return 1, "未找到有效数据区域"

    header_row = data[1] if len(data) >= 2 else []
    sub_row_indexes = [(name, rows[name] - 1) for name in SUB_REGIONS if name in rows]

    errors: list[str] = []
    for col in range(start_col, last_col + 1):
        ci = col - 1
        total_val = _to_num_or_none(total_row_data[ci] if ci < len(total_row_data) else None)
        if total_val is None:
            continue

        sub_sum = 0.0
        any_sub_present = False
        for _, ri in sub_row_indexes:
            row_data = data[ri]
            v = _to_num_or_none(row_data[ci] if ci < len(row_data) else None)
            if v is None:
                continue
            sub_sum += v
            any_sub_present = True

        if not any_sub_present:
            continue

        diff = total_val - sub_sum
        if abs(diff) > ABS_TOLERANCE:
            significant = True
        elif total_val != 0:
            significant = abs(diff / total_val) > REL_TOLERANCE
        else:
            significant = sub_sum != 0

        if significant:
            header = header_row[ci] if ci < len(header_row) else None
            header_name = str(header).strip() if header is not None else ""
            if not header_name:
                header_name = f"列{get_column_letter(col)}"
            errors.append(
                f"{header_name}: {total_val:.2f} ≠ {sub_sum:.2f} (差{diff:.2f})"
            )

    missing = [n for n in SUB_REGIONS if n not in rows]
    detail = "; ".join(errors)
    if missing:
        detail += ("; " if detail else "") + "未找到: " + ", ".join(missing)
    return len(errors), detail


def _write_result_header(ws_out) -> None:
    for i, h in enumerate(RESULT_HEADERS, start=1):
        c = ws_out.cell(row=1, column=i, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
        c.fill = HEADER_FILL


def _write_result_row(ws_out, row: int, idx: int, file_name: str, sheet_name: str,
                      err_count: int, err_detail: str) -> None:
    status = "错误" if err_count > 0 else "正确"
    detail = err_detail if err_count > 0 else "核对通过"
    values = [idx, file_name, sheet_name, status, err_count, detail]
    for j, v in enumerate(values, start=1):
        ws_out.cell(row=row, column=j, value=v)
    fill = ERROR_FILL if err_count > 0 else OK_FILL
    font = ERROR_FONT if err_count > 0 else OK_FONT
    for j in range(1, len(values) + 1):
        c = ws_out.cell(row=row, column=j)
        c.fill = fill
        c.font = font


def _autofit(ws_out) -> None:
    for col_idx in range(1, ws_out.max_column + 1):
        max_len = 0
        for row_idx in range(1, ws_out.max_row + 1):
            v = ws_out.cell(row=row_idx, column=col_idx).value
            if v is None:
                continue
            s = str(v)
            ln = sum(2 if ord(ch) > 127 else 1 for ch in s)
            if ln > max_len:
                max_len = ln
        ws_out.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 8), 80)


def run_region_sum_check(target_wbs: Iterable[Path], output_dir: Path,
                         log_dir: Path) -> dict:
    logger = get_logger("feature_1_5_region_sum_check", log_dir)

    out_wb = Workbook()
    ws_out = out_wb.active
    ws_out.title = "核对结果"
    _write_result_header(ws_out)

    total = {"files": 0, "sheets": 0, "errors": 0, "failed_files": 0, "path": None}
    out_row = 2
    for src in list(target_wbs):
        try:
            wb = load_workbook(src, data_only=True, read_only=True)
        except Exception as e:
            total["failed_files"] += 1
            logger.error(f"打开失败: {src} - {e}")
            continue

        total["files"] += 1
        file_sheets = 0
        file_errors = 0
        try:
            for ws in wb.worksheets:
                if REGION_SHEET_KEY not in ws.title:
                    continue
                err_count, err_detail = _check_one_sheet(ws)
                total["sheets"] += 1
                total["errors"] += err_count
                file_sheets += 1
                file_errors += err_count
                _write_result_row(ws_out, out_row, out_row - 1, src.name, ws.title,
                                  err_count, err_detail)
                logger.debug(f"  {ws.title}: {'有误 '+str(err_count) if err_count else '通过'}")
                out_row += 1
        finally:
            wb.close()
        logger.info(f"{src.name}: 校验 {file_sheets} 个分地区 sheet，错误 {file_errors}")

    if out_row > 2:
        ws_out.cell(row=out_row, column=1, value="=== 统计汇总 ===")
        ws_out.cell(row=out_row, column=2, value=f"已检查文件数: {total['files']}")
        ws_out.cell(row=out_row, column=3, value=f"已检查工作表数: {total['sheets']}")
        ws_out.cell(row=out_row, column=4, value=f"发现错误总数: {total['errors']}")

    _autofit(ws_out)
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / f"地区数值核对结果_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    out_wb.save(out_path)
    out_wb.close()
    total["path"] = str(out_path)
    logger.info(f"结果已保存: {out_path}")
    return total
