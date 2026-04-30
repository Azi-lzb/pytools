"""问卷调研统计 — 配置驱动版

读取一份"按机构 × 题汇总"的 Excel（典型来源：3.2.2 按批注汇总输出），
按 config.xlsx 中 `问卷统计配置` sheet 的参数生成两份报告：
  A) 明细统计版：每题每选项一行（小计/比例/作答机构数等）
  B) 展示核查版：每题分块，含选项 -> 机构名单、多选机构、未选机构

机构名称标准化复用 config.xlsx 已有的 `重命名配置`（A 列简称 → B 列全称）。
"""
from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, column_index_from_string

from .logger import get_logger


# ============== 配置数据结构 ==============

CONFIG_SHEET = "问卷统计配置"
DEFAULT_RENAME_SHEET = "重命名配置"


@dataclass
class SurveyConfig:
    name: str
    sheet_name: str            # 输入工作表名，留空 = 第一张
    org_col: str               # 机构列（列名 / 字母 / 列号）
    question_col: str
    options_col: str
    answer_col: str
    path_cols: list[str] = field(default_factory=list)
    group_mode: str = "路径+题干分组"        # 题干分组 / 路径+题干分组
    parse_mode: str = "编码文本混合"          # 编码优先 / 文本优先 / 编码文本混合
    answer_sep_regex: str = r"[,，;；/、\s]+"
    org_fallback_regex: list[str] = field(default_factory=list)
    rename_cfg_path: Path | None = None       # 留空 = 用 config.xlsx 自身
    rename_cfg_sheet: str = DEFAULT_RENAME_SHEET
    rename_src_col: str = "A"
    rename_dst_col: str = "B"
    output_prefix: str = "问卷汇总"
    output_dir: Path = Path("output")
    enabled: bool = True


# ============== 配置加载 ==============

_TRUE_TOKENS = {"是", "y", "Y", "yes", "true", "TRUE", "1", 1, True}


def _is_truthy(v) -> bool:
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    return str(v).strip() in {"是", "y", "Y", "yes", "Yes", "YES", "true", "TRUE", "True", "1"}


def _split_list(v) -> list[str]:
    """常规列表字段：用 , ， ； ; | 分隔。"""
    if v is None:
        return []
    s = str(v).strip()
    if not s:
        return []
    return [p.strip() for p in re.split(r"[,，；;|]", s) if p.strip()]


def _split_regex_list(v) -> list[str]:
    """正则列表字段：仅用 || 分隔（避免与正则元字符冲突）。"""
    if v is None:
        return []
    s = str(v).strip()
    if not s:
        return []
    return [p.strip() for p in s.split("||") if p.strip()]


def load_survey_config(config_xlsx: Path) -> SurveyConfig:
    """从 config.xlsx 的 `问卷统计配置` sheet 读取首条启用配置。"""
    wb = load_workbook(config_xlsx, data_only=True)
    if CONFIG_SHEET not in wb.sheetnames:
        raise ValueError(f"配置文件缺少工作表 `{CONFIG_SHEET}`：{config_xlsx}")
    ws = wb[CONFIG_SHEET]
    headers = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c).value
        if h is not None:
            headers[str(h).strip()] = c
    required = ["是否启用", "配置名称", "机构列", "题目列", "选项列", "答案列"]
    missing = [h for h in required if h not in headers]
    if missing:
        raise ValueError(f"`{CONFIG_SHEET}` 缺少字段：{missing}")

    def get(row: int, key: str, default=None):
        c = headers.get(key)
        if c is None:
            return default
        v = ws.cell(row, c).value
        return v if v not in (None, "") else default

    for r in range(2, ws.max_row + 1):
        if not _is_truthy(get(r, "是否启用")):
            continue
        cfg_name = str(get(r, "配置名称") or "").strip()
        if not cfg_name:
            continue

        rename_path_raw = get(r, "重命名配置文件路径")
        rename_path = Path(str(rename_path_raw)) if rename_path_raw else None

        out_dir_raw = get(r, "输出目录") or "output"
        out_dir = Path(str(out_dir_raw))

        cfg = SurveyConfig(
            name=cfg_name,
            sheet_name=str(get(r, "工作表名称") or "").strip(),
            org_col=str(get(r, "机构列") or "").strip(),
            question_col=str(get(r, "题目列") or "").strip(),
            options_col=str(get(r, "选项列") or "").strip(),
            answer_col=str(get(r, "答案列") or "").strip(),
            path_cols=_split_list(get(r, "路径列列表")),
            group_mode=str(get(r, "题目分组方式") or "路径+题干分组").strip(),
            parse_mode=str(get(r, "选项解析方式") or "编码文本混合").strip(),
            answer_sep_regex=str(get(r, "答案分隔符正则") or r"[,，;；/、\s]+"),
            org_fallback_regex=_split_regex_list(get(r, "机构名称兜底清洗规则")),
            rename_cfg_path=rename_path,
            rename_cfg_sheet=str(get(r, "重命名配置工作表") or DEFAULT_RENAME_SHEET).strip(),
            rename_src_col=str(get(r, "重命名原名列") or "A").strip(),
            rename_dst_col=str(get(r, "重命名标准名列") or "B").strip(),
            output_prefix=str(get(r, "输出前缀") or "问卷汇总").strip(),
            output_dir=out_dir,
            enabled=True,
        )
        wb.close()
        return cfg
    wb.close()
    raise ValueError(f"`{CONFIG_SHEET}` 中没有 `是否启用=是` 的配置项")


# ============== 重命名映射 ==============

def _col_to_index(col: str) -> int:
    """支持 'A'/'B'... 字母列名或纯数字。"""
    s = str(col).strip()
    if s.isdigit():
        return int(s)
    return column_index_from_string(s.upper())


def load_rename_map(cfg: SurveyConfig, default_xlsx: Path) -> list[tuple[str, str]]:
    """返回 [(简称, 全称)]，按简称长度降序排序（长匹配优先）。"""
    path = cfg.rename_cfg_path or default_xlsx
    if not Path(path).exists():
        raise FileNotFoundError(f"重命名配置文件不存在：{path}")
    wb = load_workbook(path, data_only=True)
    if cfg.rename_cfg_sheet not in wb.sheetnames:
        raise ValueError(f"重命名配置缺少工作表 `{cfg.rename_cfg_sheet}`：{path}")
    ws = wb[cfg.rename_cfg_sheet]
    sc = _col_to_index(cfg.rename_src_col)
    dc = _col_to_index(cfg.rename_dst_col)
    out: list[tuple[str, str]] = []
    for r in range(2, ws.max_row + 1):
        a, b = ws.cell(r, sc).value, ws.cell(r, dc).value
        if a is None or b is None:
            continue
        a_s, b_s = str(a).strip(), str(b).strip()
        if not a_s or not b_s:
            continue
        out.append((a_s, b_s))
    wb.close()
    out.sort(key=lambda x: -len(x[0]))
    return out


# ============== 机构名称标准化 ==============

_BASE_STRIP_RE = re.compile(r"\.(xlsx?|xlsm|csv)$", flags=re.IGNORECASE)


def _basic_clean(raw: str) -> str:
    s = (raw or "").strip()
    s = _BASE_STRIP_RE.sub("", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip()


_SHORT_SRC_THRESHOLD = 3   # 简称长度 < 此值时仅做精确匹配，避免如"广州"误命中"广州港"


def _apply_fallback(s: str, rules: list[str]) -> str:
    cur = s
    for pat in rules:
        try:
            cur = re.sub(pat, "", cur)
        except re.error:
            continue
    cur = re.sub(r"\s+", " ", cur).strip(" -_:：·、，,。.+()（）")
    return cur


def _match_rename(cand: str, rename_map: list[tuple[str, str]]) -> str | None:
    """优先精确匹配；长简称(>=阈值)再做包含匹配；短简称只允许精确。"""
    if not cand:
        return None
    # 精确（含等于 dst 的身份行）
    for src, dst in rename_map:
        if cand == src or cand == dst:
            return dst
    # 包含（rename_map 已按 src 长度降序）
    for src, dst in rename_map:
        if len(src) >= _SHORT_SRC_THRESHOLD and src in cand:
            return dst
    return None


def normalize_org_name(raw, rename_map: list[tuple[str, str]],
                       fallback_rules: list[str]) -> tuple[str, bool]:
    """
    返回 (标准名, 是否命中重命名映射)。流程：
    1) basic_clean → base
    2) 在 base 上先尝试 rename_map（精确 + 长简称包含）
    3) 否则用 fallback 正则剥噪音 → cleaned，再尝试 rename_map
    4) 仍未命中则返回 cleaned（或 base）
    """
    base = _basic_clean(str(raw or ""))
    if not base:
        return "", False

    hit = _match_rename(base, rename_map)
    if hit is not None:
        return hit, True

    cleaned = _apply_fallback(base, fallback_rules)
    hit = _match_rename(cleaned, rename_map)
    if hit is not None:
        return hit, True

    return (cleaned or base), False


# ============== 选项与答案解析 ==============

# A 城市商业银行 / B 农村商业银行  / B. xxx / B、xxx
_OPTION_LINE_RE = re.compile(r"^\s*([A-Za-z])[\s\.\、\:：\)）]\s*(.*?)\s*$")


def parse_options_text(opt_text) -> list[tuple[str, str]]:
    """解析多行选项定义，返回 [(code, text)]。"""
    if not opt_text:
        return []
    out: list[tuple[str, str]] = []
    seen: set[str] = set()
    for line in str(opt_text).splitlines():
        line = line.strip()
        if not line:
            continue
        m = _OPTION_LINE_RE.match(line)
        if m:
            code = m.group(1).upper()
            text = m.group(2).strip()
            if code in seen:
                continue
            seen.add(code)
            out.append((code, text))
    return out


def detect_question_type(question: str, opt_text: str) -> str:
    """启发式判断单选/多选：题干或选项文本中含'多选'即视为多选。"""
    blob = f"{question or ''} {opt_text or ''}"
    if "多选" in blob:
        return "多选"
    return "单选"


def parse_answer(ans, options: list[tuple[str, str]],
                 sep_regex: str, parse_mode: str) -> list[str]:
    """从答案文本中识别勾选了哪些选项编码。"""
    if ans is None:
        return []
    s = str(ans).strip()
    if not s:
        return []
    codes = [c for c, _ in options]
    code_set = set(codes)
    text_map = {t.strip(): c for c, t in options if t and t.strip()}

    raw_parts = re.split(sep_regex, s) if sep_regex else [s]
    parts = [p.strip() for p in raw_parts if p and p.strip()]
    found: set[str] = set()

    use_code = parse_mode in ("编码优先", "编码文本混合")
    use_text = parse_mode in ("文本优先", "编码文本混合")

    for p in parts:
        if use_code:
            up = p.upper()
            if up in code_set:
                found.add(up)
                continue
            m = re.match(r"^([A-Z])(?=$|[\s\.\、\:：\)）])", up)
            if m and m.group(1) in code_set:
                found.add(m.group(1))
                continue
            # 'AB' / 'ABC' 紧凑：所有字符都是有效编码且长度 <=6
            if 2 <= len(up) <= 6 and all(ch in code_set for ch in up):
                for ch in up:
                    found.add(ch)
                continue
        if use_text:
            hit = False
            if p in text_map:
                found.add(text_map[p])
                hit = True
            else:
                for code, text in options:
                    if text and (p == text or (len(p) >= 2 and p in text)):
                        found.add(code)
                        hit = True
                        break
            if hit:
                continue
    return sorted(found)


# ============== 列定位 ==============

def _resolve_col(headers: dict[str, int], spec: str) -> int | None:
    """spec 可以是列名、字母列号(A/B/...)、纯数字列号。返回 1-based 列号。"""
    if not spec:
        return None
    s = str(spec).strip()
    if s in headers:
        return headers[s]
    if s.isdigit():
        return int(s)
    if re.fullmatch(r"[A-Za-z]{1,3}", s):
        return column_index_from_string(s.upper())
    return None


# ============== 聚合统计 ==============

@dataclass
class QuestionStat:
    title: str                         # 完整题目（路径+题干 或 仅题干）
    question: str                      # 题干
    qtype: str                         # 单选 / 多选
    options: list[tuple[str, str]]     # [(code, text)]
    counts: dict[str, int] = field(default_factory=dict)   # code -> 小计
    org_lists: dict[str, list[str]] = field(default_factory=dict)  # code -> [机构]
    answered_orgs: set[str] = field(default_factory=set)
    multi_orgs: list[tuple[str, int]] = field(default_factory=list)  # 单选题里勾了多个的机构
    raw_per_org: dict[str, list[str]] = field(default_factory=dict)  # 机构 -> 该题勾的 codes


def aggregate(input_path: Path, cfg: SurveyConfig, rename_map, log) -> tuple[list[QuestionStat], set[str], int, int]:
    wb = load_workbook(input_path, data_only=True)
    ws = wb[cfg.sheet_name] if cfg.sheet_name and cfg.sheet_name in wb.sheetnames else wb.worksheets[0]
    headers: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v is not None:
            headers[str(v).strip()] = c

    org_c = _resolve_col(headers, cfg.org_col)
    q_c = _resolve_col(headers, cfg.question_col)
    opt_c = _resolve_col(headers, cfg.options_col)
    ans_c = _resolve_col(headers, cfg.answer_col)
    missing = [n for n, v in [
        ("机构列", org_c), ("题目列", q_c), ("选项列", opt_c), ("答案列", ans_c)] if v is None]
    if missing:
        wb.close()
        raise ValueError(f"输入文件中找不到列：{missing}（输入表头={list(headers.keys())}）")

    path_cs = [c for c in (_resolve_col(headers, p) for p in cfg.path_cols) if c is not None]

    use_path = cfg.group_mode.startswith("路径")
    questions: dict[str, QuestionStat] = {}  # question_text -> QuestionStat（先按题干分组）
    best_path: dict[str, list[str]] = {}     # question_text -> 最完整的路径
    sample_orgs: set[str] = set()
    rename_hits = 0
    rename_miss = 0

    for r in range(2, ws.max_row + 1):
        raw_org = ws.cell(r, org_c).value
        if raw_org is None:
            continue
        org, hit = normalize_org_name(raw_org, rename_map, cfg.org_fallback_regex)
        if not org:
            continue
        if hit:
            rename_hits += 1
        else:
            rename_miss += 1
        sample_orgs.add(org)

        question = str(ws.cell(r, q_c).value or "").strip()
        if not question:
            continue
        opt_text = ws.cell(r, opt_c).value
        ans = ws.cell(r, ans_c).value

        # 分组键统一用题干（避免同题因路径列偶有缺失被拆成多组）
        title_key = question
        if use_path and path_cs:
            path_parts = [str(ws.cell(r, c).value).strip()
                          for c in path_cs if ws.cell(r, c).value not in (None, "")]
            prev = best_path.get(question, [])
            if len(path_parts) > len(prev):
                best_path[question] = path_parts

        st = questions.get(title_key)
        if st is None:
            options = parse_options_text(opt_text)
            qtype = detect_question_type(question, opt_text or "")
            st = QuestionStat(title=question, question=question, qtype=qtype, options=options)
            for code, _ in options:
                st.counts[code] = 0
                st.org_lists[code] = []
            questions[title_key] = st
        elif not st.options:
            # 之前只见到空选项的行，这次有定义就补上
            st.options = parse_options_text(opt_text)
            for code, _ in st.options:
                st.counts.setdefault(code, 0)
                st.org_lists.setdefault(code, [])

        ticked = parse_answer(ans, st.options, cfg.answer_sep_regex, cfg.parse_mode)
        if not ticked:
            continue
        st.answered_orgs.add(org)
        prev = st.raw_per_org.setdefault(org, [])
        for code in ticked:
            if code not in st.counts:
                # 答案出现了选项里没定义的编码，跳过
                continue
            if org not in st.org_lists[code]:
                st.org_lists[code].append(org)
                st.counts[code] = st.counts.get(code, 0) + 1
            if code not in prev:
                prev.append(code)

    wb.close()

    # 用最完整的路径回填显示标题
    if use_path:
        for q_text, st in questions.items():
            parts = best_path.get(q_text, [])
            st.title = f"{'_'.join(parts)} | {q_text}" if parts else q_text

    # 过滤掉没有选项定义的题（如联系人/数据基础是否一类无选项问答），统计无意义
    skipped_no_options = [st.question for st in questions.values() if not st.options]
    if skipped_no_options:
        log.info(f"跳过无选项题数={len(skipped_no_options)}：{skipped_no_options}")

    # 标记单选题里勾了多个的机构（口径：实际勾了 >1 个不同选项）
    stat_list = [st for st in questions.values() if st.options]
    for st in stat_list:
        if st.qtype != "单选":
            continue
        for org, codes in st.raw_per_org.items():
            if len(codes) > 1:
                st.multi_orgs.append((org, len(codes)))

    log.info(f"识别题数={len(stat_list)} 样本机构数={len(sample_orgs)}")
    log.info(f"重命名命中={rename_hits} 未命中={rename_miss}")
    return stat_list, sample_orgs, rename_hits, rename_miss


# ============== 输出 A：明细统计版 ==============

def _bold(cell):
    cell.font = Font(bold=True)


def write_detail(stats: list[QuestionStat], sample_size: int,
                 cfg: SurveyConfig) -> Path:
    out = Workbook()
    ws = out.active
    ws.title = "选项比例统计"
    headers = ["题目(主路径+问题)", "题干", "题型", "选项编码", "选项文本",
               "小计", "占作答比例", "占勾选比例",
               "作答机构数", "勾选总次数", "样本机构数"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i)
        c.value = h
        _bold(c)
    r = 2
    for st in stats:
        answered = len(st.answered_orgs)
        total_ticks = sum(st.counts.values())
        for code, text in st.options:
            n = st.counts.get(code, 0)
            ratio_ans = (n / answered) if answered else 0
            ratio_tick = (n / total_ticks) if total_ticks else 0
            ws.cell(r, 1).value = st.title
            ws.cell(r, 2).value = st.question
            ws.cell(r, 3).value = st.qtype
            ws.cell(r, 4).value = code
            ws.cell(r, 5).value = text
            ws.cell(r, 6).value = n
            ws.cell(r, 7).value = ratio_ans
            ws.cell(r, 8).value = ratio_tick
            ws.cell(r, 9).value = answered
            ws.cell(r, 10).value = total_ticks
            ws.cell(r, 11).value = sample_size
            r += 1
    ws.freeze_panes = "A2"
    cfg.output_dir.mkdir(parents=True, exist_ok=True)
    p = cfg.output_dir / f"{cfg.output_prefix}_按题目选项比例统计_程序版.xlsx"
    out.save(p)
    out.close()
    return p


# ============== 输出 B：展示核查版 ==============

def _format_org_list(orgs: list[str]) -> str:
    return "；".join(orgs) if orgs else "（无）"


def write_show(stats: list[QuestionStat], sample_orgs: set[str],
               cfg: SurveyConfig) -> Path:
    out = Workbook()
    ws = out.active
    ws.title = "展示版_含机构核查"
    sample_size = len(sample_orgs)
    block_fill = PatternFill(start_color="FFEFD5", end_color="FFEFD5", fill_type="solid")
    sub_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    r = 1
    for idx, st in enumerate(stats, 1):
        title_cell = ws.cell(r, 1)
        title_cell.value = f"第{idx}题：{st.title} [{st.qtype}]"
        _bold(title_cell)
        title_cell.fill = block_fill
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        r += 1

        for ci, h in enumerate(["项", "小计", "比例", "机构列表（用于核查）"], 1):
            c = ws.cell(r, ci)
            c.value = h
            _bold(c)
            c.fill = sub_fill
        r += 1

        answered = len(st.answered_orgs)
        total_ticks = sum(st.counts.values())
        denom = total_ticks if st.qtype == "多选" else answered
        for code, text in st.options:
            n = st.counts.get(code, 0)
            ratio = (n / denom) if denom else 0
            ws.cell(r, 1).value = f"{code}. {text}"
            ws.cell(r, 2).value = n
            ws.cell(r, 3).value = ratio
            ws.cell(r, 4).value = _format_org_list(st.org_lists.get(code, []))
            r += 1

        # 多选机构（仅对单选题有意义；多选题保留行但置 0）
        multi_n = len(st.multi_orgs)
        ws.cell(r, 1).value = "多选机构"
        ws.cell(r, 2).value = multi_n
        ws.cell(r, 4).value = (
            "；".join(f"{name}({n})" for name, n in st.multi_orgs)
            if multi_n else "（无）"
        )
        r += 1

        # 未选机构
        unanswered = sorted(sample_orgs - st.answered_orgs)
        ws.cell(r, 1).value = "未选机构"
        ws.cell(r, 2).value = len(unanswered)
        ws.cell(r, 4).value = _format_org_list(unanswered)
        r += 1

        # 本题有效填写机构数
        ws.cell(r, 1).value = "本题有效填写机构数"
        ws.cell(r, 2).value = answered
        ws.cell(r, 4).value = f"样本机构数：{sample_size}"
        r += 2  # 空一行

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 80
    for row in ws.iter_rows(min_row=1, max_row=r):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    cfg.output_dir.mkdir(parents=True, exist_ok=True)
    p = cfg.output_dir / f"{cfg.output_prefix}_按题目选项比例统计_展示版_程序版.xlsx"
    out.save(p)
    out.close()
    return p


# ============== 入口 ==============

def run_survey_stats(input_path: Path, config_xlsx: Path,
                     log_dir: Path) -> dict:
    log = get_logger("survey_stats", log_dir)
    log.info(f"输入文件：{input_path}")
    cfg = load_survey_config(config_xlsx)
    log.info(f"使用配置：{cfg.name}")

    rename_map = load_rename_map(cfg, default_xlsx=config_xlsx)
    log.info(f"重命名映射条数：{len(rename_map)}")

    stats, sample_orgs, hits, miss = aggregate(input_path, cfg, rename_map, log)
    if not stats:
        log.warning("未识别到任何题目，请检查列映射与输入数据。")
        return {"saved": False, "questions": 0, "sample": 0,
                "rename_hits": hits, "rename_miss": miss,
                "detail": "", "show": ""}

    detail_p = write_detail(stats, len(sample_orgs), cfg)
    show_p = write_show(stats, sample_orgs, cfg)
    log.info(f"输出明细：{detail_p}")
    log.info(f"输出展示：{show_p}")
    return {
        "saved": True,
        "questions": len(stats),
        "sample": len(sample_orgs),
        "rename_hits": hits,
        "rename_miss": miss,
        "detail": str(detail_p),
        "show": str(show_p),
    }
