from __future__ import annotations

import shutil
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .logger import get_logger


BRANCH_SHEET_KEY = "分机构"
COPY_SUFFIX = "(金融局)"


def _find_last_row_colA(ws) -> int:
    for r in range(ws.max_row, 0, -1):
        v = ws.cell(row=r, column=1).value
        if v is not None and str(v).strip() != "":
            return r
    return 0


def _delete_foreign_rows(ws, foreign_set: set[str]) -> int:
    """A 列机构名命中外资行集合的整行删除。返回删除行数。"""
    last_row = _find_last_row_colA(ws)
    if last_row < 1:
        return 0

    marked: list[bool] = [False] * (last_row + 1)  # 1-based
    for r in range(1, last_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is None:
            continue
        name = str(v).strip()
        if name and name in foreign_set:
            marked[r] = True

    # 自底向上合并连续命中块批量删除
    deleted = 0
    r = last_row
    while r >= 1:
        if marked[r]:
            end = r
            while r >= 1 and marked[r]:
                r -= 1
            start = r + 1
            ws.delete_rows(start, end - start + 1)
            deleted += end - start + 1
        else:
            r -= 1
    return deleted


def _build_copy_path(src: Path) -> Path:
    return src.with_name(f"{src.stem}{COPY_SUFFIX}{src.suffix}")


def _process_one_workbook(src: Path, foreign_set: set[str], logger) -> dict:
    stat = {"branch_sheets": 0, "deleted_rows": 0, "copy_path": None, "ok": False}
    dst = _build_copy_path(src)
    shutil.copy2(src, dst)
    wb = load_workbook(dst, keep_vba=dst.suffix.lower() == ".xlsm")
    try:
        for ws in wb.worksheets:
            if BRANCH_SHEET_KEY not in ws.title:
                continue
            stat["branch_sheets"] += 1
            n = _delete_foreign_rows(ws, foreign_set)
            if n > 0:
                stat["deleted_rows"] += n
                logger.info(f"  sheet '{ws.title}' 删除 {n} 行")
        wb.save(dst)
        stat["ok"] = True
        stat["copy_path"] = str(dst)
    finally:
        wb.close()
    return stat


def run_remove_foreign_bank(target_wbs: Iterable[Path], foreign_set: set[str],
                            log_dir: Path) -> dict:
    logger = get_logger("feature_1_3_remove_foreign_bank", log_dir)
    if not foreign_set:
        logger.warning("外资行配置为空（机构映射表 C 列无 1），无法继续")
        return {"files": 0, "ok_files": 0, "branch_sheets": 0,
                "deleted_rows": 0, "failed_files": 0, "foreign_empty": True}

    logger.info(f"外资行数量: {len(foreign_set)}")

    total = {"files": 0, "ok_files": 0, "branch_sheets": 0,
             "deleted_rows": 0, "failed_files": 0, "foreign_empty": False}
    for src in target_wbs:
        total["files"] += 1
        try:
            logger.info(f"处理文件: {src}")
            stat = _process_one_workbook(src, foreign_set, logger)
            total["branch_sheets"] += stat["branch_sheets"]
            total["deleted_rows"] += stat["deleted_rows"]
            if stat["ok"]:
                total["ok_files"] += 1
                logger.info(
                    f"文件完成: {Path(stat['copy_path']).name} "
                    f"分机构={stat['branch_sheets']} 删除={stat['deleted_rows']}"
                )
        except Exception as e:
            total["failed_files"] += 1
            logger.error(f"文件失败: {src} - {e}")
    return total
