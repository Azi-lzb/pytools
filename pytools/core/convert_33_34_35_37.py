from __future__ import annotations

from pathlib import Path
import os
import re
import gc
import shutil
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string

from .config_xlsx import (
    SHEET_ARCHIVE_TYPE_CONFIG,
    SHEET_CONFIG_RENAME,
    SHEET_FILE_SORT_CONFIG,
    SHEET_SUBMISSION_CHECK_CONFIG,
)
from .io_excel import write_workbook
from .logger import get_logger


def _norm(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _split_tokens(v: str) -> list[str]:
    s = _norm(v).replace("；", ";").replace(",", ";")
    return [x.strip() for x in s.split(";") if x.strip()]


def _is_enabled(v) -> bool:
    return _norm(v).lower() in ("1", "y", "yes", "true", "是", "启用")


def _workbook_supported(path: Path) -> bool:
    return path.suffix.lower() in (".xlsx", ".xlsm")


def _excel_like(path: Path) -> bool:
    return path.suffix.lower() in (".xls", ".xlsx", ".xlsm", ".xlsb", ".xlt", ".xltx", ".xltm", ".csv")


def _load_config_rename_df(cfg_path: Path) -> pd.DataFrame:
    return pd.read_excel(cfg_path, SHEET_CONFIG_RENAME, header=0, dtype=object)


def _build_rename_maps(cfg_path: Path) -> tuple[dict[str, str], dict[str, str], dict[str, str]]:
    df = _load_config_rename_df(cfg_path)
    short_to_full: dict[str, str] = {}
    full_to_code: dict[str, str] = {}
    kv: dict[str, str] = {}

    for _, row in df.iterrows():
        a = _norm(row.get("简称"))
        b = _norm(row.get("全称"))
        d = _norm(row.get("代码"))
        e = _norm(row.get("全称(代码映射)"))
        g = _norm(row.get("键"))
        h = _norm(row.get("值"))
        if a and b:
            short_to_full[a] = b
        if e and d:
            full_to_code[e] = d
        if g and h:
            kv[g] = h
    return short_to_full, full_to_code, kv


def _next_path(p: Path) -> Path:
    if not p.exists():
        return p
    i = 1
    while True:
        cand = p.with_name(f"{p.stem}_{i}{p.suffix}")
        if not cand.exists():
            return cand
        i += 1


def _next_plan_path(p: Path, planned: set[Path]) -> Path:
    cand = p
    i = 1
    while cand.exists() or cand in planned:
        cand = p.with_name(f"{p.stem}_{i}{p.suffix}")
        i += 1
    planned.add(cand)
    return cand


def _safe_folder_name(text: str) -> str:
    s = _sanitize_name_part(text)
    return s or "未识别"


def _iter_archive_candidates(paths: list[Path]) -> tuple[list[Path], list[Path]]:
    files: list[Path] = []
    invalid: list[Path] = []
    seen: set[Path] = set()
    for p in paths:
        if p.is_dir():
            for child in sorted(p.rglob("*")):
                if child.is_file() and not child.name.startswith("~$"):
                    key = child.resolve()
                    if key not in seen:
                        files.append(child)
                        seen.add(key)
        elif p.is_file():
            key = p.resolve()
            if key not in seen and not p.name.startswith("~$"):
                files.append(p)
                seen.add(key)
        else:
            invalid.append(p)
    return files, invalid


def _parse_archive_exts(raw) -> set[str]:
    text = _norm(raw)
    if not text:
        return set()
    parts = re.split(r"[;；,，\s]+", text)
    out: set[str] = set()
    for part in parts:
        ext = part.strip().lower()
        if not ext:
            continue
        ext = ext.lstrip(".")
        if ext:
            out.add(f".{ext}")
    return out


def _load_archive_ext_filter(cfg_path: Path | None) -> tuple[set[str], set[str], str]:
    if cfg_path is None or not cfg_path.exists():
        return set(), set(), "未配置后缀筛选，按全部文件处理"
    try:
        wb = load_workbook(cfg_path, read_only=True, data_only=True)
        try:
            if SHEET_ARCHIVE_TYPE_CONFIG not in wb.sheetnames:
                return set(), set(), "未配置后缀筛选，按全部文件处理"
            ws = wb[SHEET_ARCHIVE_TYPE_CONFIG]
            include = _parse_archive_exts(ws["E2"].value)
            exclude = _parse_archive_exts(ws["F2"].value)
            return include, exclude, ""
        finally:
            wb.close()
    except Exception:
        return set(), set(), "未配置后缀筛选，按全部文件处理"


def _archive_ext_skip_status(path: Path, include_exts: set[str], exclude_exts: set[str]) -> str:
    ext = path.suffix.lower()
    if ext in exclude_exts:
        return "跳过：后缀被排除"
    if include_exts and ext not in include_exts:
        return "跳过：后缀不在归档范围"
    return ""


def _archive_date_from_name(name: str) -> str:
    text = Path(name).stem

    patterns = [
        r"(?<!\d)(20\d{2})[-_.年](0?[1-9]|1[0-2])[-_.月](0?[1-9]|[12]\d|3[01])日?(?!\d)",
        r"(?<!\d)(20\d{2})(0[1-9]|1[0-2])([0-2]\d|3[01])(?!\d)",
        r"(?<!\d)(20\d{2})[-_.年](0?[1-9]|1[0-2])月?(?!\d)",
        r"(?<!\d)(20\d{2})(0[1-9]|1[0-2])(?!\d)",
    ]
    for pattern in patterns:
        m = re.search(pattern, text)
        if not m:
            continue
        y = m.group(1)
        mo = int(m.group(2))
        if len(m.groups()) >= 3 and m.group(3) is not None:
            day = int(m.group(3))
            if 1 <= mo <= 12 and 1 <= day <= 31:
                return f"{y}{mo:02d}"
        elif 1 <= mo <= 12:
            return f"{y}{mo:02d}"
    return ""


def _load_archive_type_rules(cfg_path: Path) -> list[tuple[str, str]]:
    df = pd.read_excel(cfg_path, SHEET_ARCHIVE_TYPE_CONFIG, header=0, dtype=object)
    rules: list[tuple[str, str]] = []
    for _, row in df.iterrows():
        if not _is_enabled(row.get("是否启用")):
            continue
        keyword = _norm(row.get("匹配关键词"))
        folder = _safe_folder_name(row.get("归档文件夹"))
        if keyword:
            rules.append((keyword, folder))
    return rules


def _load_submission_check_rules(cfg_path: Path) -> list[dict[str, object]]:
    df = pd.read_excel(cfg_path, SHEET_SUBMISSION_CHECK_CONFIG, header=0, dtype=object).fillna("")
    rules: list[dict[str, object]] = []
    for idx, row in df.iterrows():
        if not _is_enabled(row.get("是否启用")):
            continue
        groups = []
        for col in ("日期关键词", "县区关键词", "机构关键词", "备用关键词"):
            groups.append(_split_tokens(row.get(col)))
        rules.append({
            "row_no": int(idx) + 2,
            "enabled": _norm(row.get("是否启用")),
            "date_keyword": _norm(row.get("日期关键词")),
            "area_keyword": _norm(row.get("县区关键词")),
            "institution_keyword": _norm(row.get("机构关键词")),
            "extra_keyword": _norm(row.get("备用关键词")),
            "remark": _norm(row.get("备注")),
            "groups": groups,
        })
    return rules


def _file_matches_submission_rule(file_name: str, rule: dict[str, object]) -> bool:
    name = file_name.lower()
    has_condition = False
    for group in rule.get("groups") or []:
        tokens = [str(x).strip().lower() for x in (group or []) if str(x).strip()]
        if not tokens:
            continue
        has_condition = True
        if not any(token in name for token in tokens):
            return False
    return has_condition


def run_submission_check_48(cfg_path: Path, files: list[Path], output_dir: Path, log_dir: Path) -> dict[str, object]:
    log = get_logger("4_8_submission_check", log_dir)
    rules = _load_submission_check_rules(cfg_path)
    if not rules:
        raise ValueError("提交检查配置 没有启用的检查项。")

    valid_files = [p for p in files if p.exists() and p.is_file()]
    file_hits: dict[Path, list[dict[str, object]]] = {p: [] for p in valid_files}
    result_rows: list[dict[str, object]] = []

    for rule in rules:
        matched = [p for p in valid_files if _file_matches_submission_rule(p.name, rule)]
        for p in matched:
            file_hits[p].append(rule)
        result_rows.append({
            "是否启用": rule["enabled"],
            "日期关键词": rule["date_keyword"],
            "县区关键词": rule["area_keyword"],
            "机构关键词": rule["institution_keyword"],
            "备用关键词": rule["extra_keyword"],
            "命中次数": len(matched),
            "命中的文件名": "；".join(p.name for p in matched),
            "备注": rule["remark"],
            "配置行": rule["row_no"],
        })

    unmatched_rows = [
        {"文件名": p.name, "源文件路径": str(p), "备注": "未命中任何启用配置"}
        for p, hits in file_hits.items()
        if not hits
    ]
    multi_rows = []
    for p, hits in file_hits.items():
        if len(hits) <= 1:
            continue
        hit_text = "；".join(
            f"第{h['row_no']}行({h['date_keyword']}|{h['area_keyword']}|{h['institution_keyword']}|{h['extra_keyword']})"
            for h in hits
        )
        multi_rows.append({"文件名": p.name, "源文件路径": str(p), "命中配置": hit_text})

    result_df = pd.DataFrame(result_rows)
    missing_df = result_df[result_df["命中次数"] == 0].copy()
    unmatched_df = pd.DataFrame(unmatched_rows, columns=["文件名", "源文件路径", "备注"])
    multi_df = pd.DataFrame(multi_rows, columns=["文件名", "源文件路径", "命中配置"])

    out_path = output_dir / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_提交缺失检查.xlsx"
    write_workbook(out_path, {
        "检查结果": result_df,
        "缺失清单": missing_df,
        "未识别文件": unmatched_df,
        "多重命中文件": multi_df,
    })

    submitted = int((result_df["命中次数"] > 0).sum()) if not result_df.empty else 0
    missing = int((result_df["命中次数"] == 0).sum()) if not result_df.empty else 0
    log.info(
        "output=%s rules=%s files=%s submitted=%s missing=%s unmatched=%s multi=%s",
        out_path, len(rules), len(valid_files), submitted, missing, len(unmatched_rows), len(multi_rows),
    )
    return {
        "output": str(out_path),
        "rules": len(rules),
        "files": len(valid_files),
        "submitted": submitted,
        "missing": missing,
        "unmatched": len(unmatched_rows),
        "multi": len(multi_rows),
    }


def _sort_prefix_digits(value) -> str:
    raw = _norm(value)
    if not raw:
        return ""
    if re.fullmatch(r"\d+(\.0+)?", raw):
        return str(int(float(raw)))
    return re.sub(r"\D", "", raw)


def _sort_prefix_width(value) -> int:
    raw = _norm(value)
    if not raw:
        return 0
    if re.fullmatch(r"\d+(\.0+)?", raw):
        raw = raw.split(".", 1)[0]
    digits = re.sub(r"\D", "", raw)
    return len(digits)


def _normalize_sort_prefix(value, width: int) -> str:
    digits = _sort_prefix_digits(value)
    if not digits:
        return ""
    return str(int(digits)).zfill(max(width, 1))


def _load_file_sort_rules(cfg_path: Path) -> list[dict[str, object]]:
    df = pd.read_excel(cfg_path, SHEET_FILE_SORT_CONFIG, header=0, dtype=object).fillna("")
    width = 1
    for _, row in df.iterrows():
        if not _is_enabled(row.get("是否启用")):
            continue
        if not _split_tokens(row.get("匹配关键词")):
            continue
        width = max(width, _sort_prefix_width(row.get("排序前缀")))
    rules: list[dict[str, object]] = []
    for idx, row in df.iterrows():
        if not _is_enabled(row.get("是否启用")):
            continue
        prefix = _normalize_sort_prefix(row.get("排序前缀"), width)
        keywords = _split_tokens(row.get("匹配关键词"))
        if not prefix or not keywords:
            continue
        rules.append({
            "row_no": int(idx) + 2,
            "prefix": prefix,
            "keywords": keywords,
            "keywords_text": _norm(row.get("匹配关键词")),
            "remark": _norm(row.get("备注")),
        })
    return rules


def _match_file_sort_rule(file_name: str, rules: list[dict[str, object]]) -> dict[str, object] | None:
    lower_name = file_name.lower()
    for rule in rules:
        for kw in rule.get("keywords") or []:
            if str(kw).lower() in lower_name:
                return rule
    return None


def _strip_sort_prefix(file_name: str) -> str:
    stem = Path(file_name).stem
    suffix = Path(file_name).suffix
    stem = re.sub(r"^\d{1,3}[_\-\s]+", "", stem)
    return f"{stem}{suffix}"


def run_file_sort_prefix_49(cfg_path: Path, files: list[Path], output_dir: Path, log_dir: Path) -> dict[str, object]:
    log = get_logger("4_9_file_sort_prefix", log_dir)
    rules = _load_file_sort_rules(cfg_path)
    if not rules:
        raise ValueError("文件排序配置 没有启用且有效的排序规则。")

    rows: list[dict[str, object]] = []
    renamed = unmatched = skipped = failed = 0
    for p in files:
        row = {
            "源文件路径": str(p),
            "原文件名": p.name,
            "新文件名": "",
            "命中配置行": "",
            "排序前缀": "",
            "匹配关键词": "",
            "状态": "",
            "备注": "",
        }
        if not p.exists() or not p.is_file():
            skipped += 1
            row["状态"] = "跳过：非文件"
            rows.append(row)
            continue

        rule = _match_file_sort_rule(p.name, rules)
        if rule is None:
            unmatched += 1
            row["状态"] = "未命中"
            rows.append(row)
            continue

        prefix = str(rule["prefix"])
        base_name = _strip_sort_prefix(p.name)
        target = p.with_name(f"{prefix}_{base_name}")
        status = "已重命名"
        try:
            if target.resolve() == p.resolve():
                row["新文件名"] = p.name
            else:
                if target.exists():
                    target = _next_path(target)
                    status = "目标同名：已自动后缀"
                os.rename(p, target)
                row["新文件名"] = target.name
            renamed += 1
            row["状态"] = status
            row["命中配置行"] = rule["row_no"]
            row["排序前缀"] = prefix
            row["匹配关键词"] = rule["keywords_text"]
            row["备注"] = rule["remark"]
            log.info("ok old=%s new=%s row=%s prefix=%s", p.name, row["新文件名"], rule["row_no"], prefix)
        except Exception as e:
            failed += 1
            row["状态"] = "失败"
            row["备注"] = str(e)
            log.warning("failed file=%s err=%s", p, e)
        rows.append(row)

    out_path = output_dir / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_文件排序前缀结果.xlsx"
    write_workbook(out_path, {"文件排序前缀结果": pd.DataFrame(rows)})
    return {
        "output": str(out_path),
        "rules": len(rules),
        "files": len(files),
        "renamed": renamed,
        "unmatched": unmatched,
        "skipped": skipped,
        "failed": failed,
    }


def _archive_type_from_name(name: str, rules: list[tuple[str, str]]) -> str:
    lower_name = name.lower()
    for keyword, folder in rules:
        if keyword.lower() in lower_name:
            return folder
    return ""


def _build_archive_preview(
    paths: list[Path],
    target_root: Path,
    mode: str,
    output_dir: Path,
    log_dir: Path,
    cfg_path: Path | None = None,
) -> dict[str, object]:
    log = get_logger(f"archive_{mode}", log_dir)
    target_root = target_root.resolve()
    files, invalid = _iter_archive_candidates(paths)
    rules = _load_archive_type_rules(cfg_path) if mode == "type" and cfg_path is not None else []
    include_exts, exclude_exts, filter_note = _load_archive_ext_filter(cfg_path)
    rows: list[dict[str, str]] = []
    planned: set[Path] = set()
    skipped_ext = 0

    for p in invalid:
        rows.append({
            "源文件路径": str(p),
            "文件名": p.name,
            "分类方式": "按日期" if mode == "date" else "按类型",
            "分类结果": "",
            "目标文件夹": "",
            "目标文件路径": "",
            "状态": "跳过：非文件",
            "备注": "",
        })

    for p in files:
        skip_status = _archive_ext_skip_status(p, include_exts, exclude_exts)
        if skip_status:
            skipped_ext += 1
            rows.append({
                "源文件路径": str(p),
                "文件名": p.name,
                "分类方式": "按日期" if mode == "date" else "按类型",
                "分类结果": "",
                "目标文件夹": "",
                "目标文件路径": "",
                "状态": skip_status,
                "备注": "",
            })
            continue

        if mode == "date":
            folder = _archive_date_from_name(p.name)
            status = "待归档" if folder else "未识别日期"
            folder = folder or "未识别日期"
        else:
            folder = _archive_type_from_name(p.name, rules)
            status = "待归档" if folder else "未识别类型"
            folder = folder or "未识别类型"

        folder = _safe_folder_name(folder)
        target_dir = target_root / folder
        raw_target = target_dir / p.name
        target = _next_plan_path(raw_target, planned)
        note = "目标同名：将自动后缀" if target.name != p.name else ""
        rows.append({
            "源文件路径": str(p),
            "文件名": p.name,
            "分类方式": "按日期" if mode == "date" else "按类型",
            "分类结果": folder,
            "目标文件夹": str(target_dir),
            "目标文件路径": str(target),
            "状态": status,
            "备注": note,
        })

    df = pd.DataFrame(rows, columns=[
        "源文件路径", "文件名", "分类方式", "分类结果",
        "目标文件夹", "目标文件路径", "状态", "备注",
    ])
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_归档计划.xlsx"
    write_workbook(out_path, {"归档计划": df})
    log.info(
        "archive preview mode=%s inputs=%s files=%s skipped_ext=%s include_exts=%s exclude_exts=%s output=%s",
        mode, len(paths), len(files), skipped_ext, sorted(include_exts), sorted(exclude_exts), out_path,
    )
    return {
        "files": len(files),
        "invalid": len(invalid),
        "planned": len(rows),
        "rules": len(rules),
        "skipped_ext": skipped_ext,
        "include_exts": sorted(include_exts),
        "exclude_exts": sorted(exclude_exts),
        "filter_note": filter_note,
        "target_root": str(target_root),
        "output": out_path,
    }


def _copy_archive_rows(rows: list[dict[str, str]], output_dir: Path, log_dir: Path) -> dict[str, object]:
    log = get_logger("archive_copy", log_dir)
    result_rows: list[dict[str, str]] = []
    copied = skip = failed = 0
    executable = {"待归档", "未识别日期", "未识别类型"}
    planned: set[Path] = set()

    for row in rows:
        src = Path(_norm(row.get("源文件路径")))
        target_raw = _norm(row.get("目标文件路径"))
        plan_status = _norm(row.get("状态"))
        result_status = ""
        actual_target = ""
        note = ""

        if plan_status not in executable:
            skip += 1
            result_status = "跳过：计划状态不可执行"
            note = plan_status
        elif not target_raw:
            skip += 1
            result_status = "跳过：目标路径为空"
        elif not src.exists() or not src.is_file():
            skip += 1
            result_status = "跳过：源文件不存在"
        else:
            try:
                target = _next_plan_path(Path(target_raw), planned)
                target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(src, target)
                copied += 1
                actual_target = str(target)
                result_status = "已复制"
                if str(target) != target_raw:
                    note = "目标同名：已自动后缀"
                log.info("copied src=%s target=%s", src, target)
            except Exception as e:
                failed += 1
                result_status = f"失败：{e}"
                log.warning("copy failed src=%s target=%s err=%s", src, target_raw, e)

        result_rows.append({
            "源文件路径": str(src),
            "原计划目标路径": target_raw,
            "实际目标路径": actual_target,
            "执行状态": result_status,
            "备注": note,
        })

    out_df = pd.DataFrame(result_rows, columns=["源文件路径", "原计划目标路径", "实际目标路径", "执行状态", "备注"])
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_归档执行结果.xlsx"
    write_workbook(out_path, {"归档执行结果": out_df})
    return {
        "rows": len(result_rows),
        "copied": copied,
        "skip": skip,
        "failed": failed,
        "output": out_path,
    }


def _build_archive_rows(
    paths: list[Path],
    target_root: Path,
    mode: str,
    log_dir: Path,
    cfg_path: Path | None = None,
) -> tuple[list[dict[str, str]], dict[str, object]]:
    log = get_logger(f"archive_{mode}", log_dir)
    target_root = target_root.resolve()
    files, invalid = _iter_archive_candidates(paths)
    rules = _load_archive_type_rules(cfg_path) if mode == "type" and cfg_path is not None else []
    include_exts, exclude_exts, filter_note = _load_archive_ext_filter(cfg_path)
    rows: list[dict[str, str]] = []
    planned: set[Path] = set()
    skipped_ext = 0

    for p in invalid:
        rows.append({
            "源文件路径": str(p),
            "文件名": p.name,
            "分类方式": "按日期" if mode == "date" else "按类型",
            "分类结果": "",
            "目标文件夹": "",
            "目标文件路径": "",
            "状态": "跳过：非文件",
            "备注": "",
        })

    for p in files:
        skip_status = _archive_ext_skip_status(p, include_exts, exclude_exts)
        if skip_status:
            skipped_ext += 1
            rows.append({
                "源文件路径": str(p),
                "文件名": p.name,
                "分类方式": "按日期" if mode == "date" else "按类型",
                "分类结果": "",
                "目标文件夹": "",
                "目标文件路径": "",
                "状态": skip_status,
                "备注": "",
            })
            continue

        if mode == "date":
            folder = _archive_date_from_name(p.name)
            status = "待归档" if folder else "未识别日期"
            folder = folder or "未识别日期"
        else:
            folder = _archive_type_from_name(p.name, rules)
            status = "待归档" if folder else "未识别类型"
            folder = folder or "未识别类型"

        folder = _safe_folder_name(folder)
        target_dir = target_root / folder
        raw_target = target_dir / p.name
        target = _next_plan_path(raw_target, planned)
        note = "目标同名：将自动后缀" if target.name != p.name else ""
        rows.append({
            "源文件路径": str(p),
            "文件名": p.name,
            "分类方式": "按日期" if mode == "date" else "按类型",
            "分类结果": folder,
            "目标文件夹": str(target_dir),
            "目标文件路径": str(target),
            "状态": status,
            "备注": note,
        })

    log.info(
        "archive rows mode=%s inputs=%s files=%s skipped_ext=%s include_exts=%s exclude_exts=%s",
        mode, len(paths), len(files), skipped_ext, sorted(include_exts), sorted(exclude_exts),
    )
    return rows, {
        "files": len(files),
        "invalid": len(invalid),
        "planned": len(rows),
        "rules": len(rules),
        "skipped_ext": skipped_ext,
        "include_exts": sorted(include_exts),
        "exclude_exts": sorted(exclude_exts),
        "filter_note": filter_note,
        "target_root": str(target_root),
    }


def run_archive_by_date_preview(
    paths: list[Path],
    target_root: Path,
    output_dir: Path,
    log_dir: Path,
    cfg_path: Path | None = None,
) -> dict[str, object]:
    return _build_archive_preview(paths, target_root, "date", output_dir, log_dir, cfg_path=cfg_path)


def run_archive_by_type_preview(
    cfg_path: Path,
    paths: list[Path],
    target_root: Path,
    output_dir: Path,
    log_dir: Path,
) -> dict[str, object]:
    return _build_archive_preview(paths, target_root, "type", output_dir, log_dir, cfg_path=cfg_path)


def run_archive_by_date_copy(
    paths: list[Path],
    target_root: Path,
    output_dir: Path,
    log_dir: Path,
    cfg_path: Path | None = None,
) -> dict[str, object]:
    rows, meta = _build_archive_rows(paths, target_root, "date", log_dir, cfg_path=cfg_path)
    stat = _copy_archive_rows(rows, output_dir, log_dir)
    stat.update(meta)
    return stat


def run_archive_by_type_copy(
    cfg_path: Path,
    paths: list[Path],
    target_root: Path,
    output_dir: Path,
    log_dir: Path,
) -> dict[str, object]:
    rows, meta = _build_archive_rows(paths, target_root, "type", log_dir, cfg_path=cfg_path)
    stat = _copy_archive_rows(rows, output_dir, log_dir)
    stat.update(meta)
    return stat


def run_archive_plan_copy(plan_path: Path, output_dir: Path, log_dir: Path) -> dict[str, object]:
    log = get_logger("archive_copy_plan", log_dir)
    df = pd.read_excel(plan_path, sheet_name="归档计划", dtype=object).fillna("")
    required = ["源文件路径", "目标文件路径", "状态"]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"归档计划缺少必要列: {', '.join(missing)}")

    rows: list[dict[str, str]] = []
    copied = skip = failed = 0
    executable = {"待归档", "未识别日期", "未识别类型"}
    planned: set[Path] = set()

    for _, row in df.iterrows():
        src = Path(_norm(row.get("源文件路径")))
        target_raw = _norm(row.get("目标文件路径"))
        plan_status = _norm(row.get("状态"))
        result_status = ""
        actual_target = ""
        note = ""

        if plan_status not in executable:
            skip += 1
            result_status = "跳过：计划状态不可执行"
            note = plan_status
        elif not target_raw:
            skip += 1
            result_status = "跳过：目标路径为空"
        elif not src.exists() or not src.is_file():
            skip += 1
            result_status = "跳过：源文件不存在"
        else:
            try:
                target = _next_plan_path(Path(target_raw), planned)
                target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(src, target)
                copied += 1
                actual_target = str(target)
                result_status = "已复制"
                if str(target) != target_raw:
                    note = "目标同名：已自动后缀"
                log.info("copied src=%s target=%s", src, target)
            except Exception as e:
                failed += 1
                result_status = f"失败：{e}"
                log.warning("copy failed src=%s target=%s err=%s", src, target_raw, e)

        rows.append({
            "源文件路径": str(src),
            "原计划目标路径": target_raw,
            "实际目标路径": actual_target,
            "执行状态": result_status,
            "备注": note,
        })

    out_df = pd.DataFrame(rows, columns=["源文件路径", "原计划目标路径", "实际目标路径", "执行状态", "备注"])
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_归档执行结果.xlsx"
    write_workbook(out_path, {"归档执行结果": out_df})
    return {
        "rows": len(rows),
        "copied": copied,
        "skip": skip,
        "failed": failed,
        "output": out_path,
    }


def _dispatch_first(progids: list[str]):
    try:
        import win32com.client  # type: ignore
    except Exception as e:
        raise RuntimeError(f"缺少 pywin32，无法执行 COM 转换: {e}")
    last_err = None
    for pid in progids:
        try:
            app = win32com.client.DispatchEx(pid)
            return app, pid
        except Exception as e:
            last_err = e
            continue
    raise RuntimeError(f"无法创建 COM 应用（已尝试: {', '.join(progids)}）: {last_err}")


def run_batch_rename_33(cfg_path: Path, files: list[Path], log_dir: Path) -> dict[str, int]:
    log = get_logger("3_3_batch_rename", log_dir)
    short_to_full, full_to_code, kv = _build_rename_maps(cfg_path)
    date_text = kv.get("数据日期", "")
    report_name = kv.get("报表名称", "")
    if not short_to_full or not full_to_code or not date_text or not report_name:
        raise ValueError("重命名配置 缺少必要配置：简称/全称、全称(代码映射)/代码、数据日期、报表名称。")

    ok = skip = 0
    for p in files:
        if not p.exists():
            skip += 1
            log.warning("skip file=%s reason=not_exists", p)
            continue
        name = p.name
        matched_short = ""
        for s in short_to_full.keys():
            if s.lower() in name.lower():
                matched_short = s
                break
        if not matched_short:
            skip += 1
            log.info("skip file=%s reason=short_name_not_matched", name)
            continue
        full_name = short_to_full.get(matched_short, "")
        code = full_to_code.get(full_name, "")
        if not code:
            skip += 1
            log.info("skip file=%s reason=full_name_no_code full=%s", name, full_name)
            continue

        new_name = f"{date_text} {code} {full_name} {report_name}{p.suffix}"
        tgt = p.with_name(new_name)
        if tgt.resolve() == p.resolve():
            skip += 1
            log.info("skip file=%s reason=name_same", name)
            continue
        if tgt.exists():
            tgt = _next_path(tgt)
        try:
            os.rename(p, tgt)
            ok += 1
            log.info("ok old=%s new=%s", name, tgt.name)
        except Exception as e:
            skip += 1
            log.warning("skip file=%s reason=rename_failed err=%s", name, e)
    return {"ok": ok, "skip": skip}


_EXCEL_FMT_MAP = {
    "xls": 56, "xlsx": 51, "xlsm": 52, "csv": 6, "xlt": 17, "xltx": 54, "xltm": 53, "xlsb": 50,
}


def _start_excel_app():
    app, engine = _dispatch_first([
        "Excel.Application",      # Microsoft Excel
        "ket.Application",        # WPS 表格（常见）
        "KET.Application",        # WPS 表格（大小写变体）
    ])
    app.Visible = False
    app.DisplayAlerts = False
    for attr, val in (("ScreenUpdating", False), ("EnableEvents", False),
                      ("AskToUpdateLinks", False), ("AlertBeforeOverwriting", False)):
        try:
            setattr(app, attr, val)
        except Exception:
            pass
    try:
        app.Calculation = -4135  # xlCalculationManual
    except Exception:
        pass
    return app, engine


def _excel_convert_one(app, src: Path, out: Path, target_ext: str) -> None:
    wb = app.Workbooks.Open(str(src), UpdateLinks=0, ReadOnly=True)
    try:
        wb.SaveAs(str(out), FileFormat=_EXCEL_FMT_MAP[target_ext], CreateBackup=False)
    finally:
        wb.Close(SaveChanges=False)


def run_excel_convert_34(
    files: list[Path], target_format: str, log_dir: Path
) -> dict[str, int]:
    log = get_logger("3_4_excel_convert", log_dir)
    ext = _norm(target_format).lower().lstrip(".") or "xlsx"
    if ext not in _EXCEL_FMT_MAP:
        ext = "xlsx"

    # 预筛有效文件，避免为 0 个有效文件也启动 Excel
    valid: list[tuple[Path, Path]] = []
    ok = skip = 0
    for p in files:
        if not p.exists() or not _excel_like(p):
            skip += 1
            log.warning("skip file=%s reason=not_excel_like", p)
            continue
        out_dir = p.parent / ext
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = out_dir / f"{p.stem}.{ext}"
        if out_path.resolve() == p.resolve():
            skip += 1
            log.info("skip file=%s reason=source_equals_target", p.name)
            continue
        if out_path.exists():
            out_path = _next_path(out_path)
        valid.append((p, out_path))

    if not valid:
        return {"ok": 0, "skip": skip, "target_ext": ext, "engine": ""}

    app, engine = _start_excel_app()
    try:
        for src, out_path in valid:
            try:
                _excel_convert_one(app, src, out_path, ext)
                ok += 1
                log.info("ok src=%s out=%s engine=%s", src, out_path, engine)
            except Exception as e:
                skip += 1
                log.warning("skip file=%s reason=convert_failed err=%s", src, e)
    finally:
        try:
            app.Quit()
        except Exception:
            pass
    return {"ok": ok, "skip": skip, "target_ext": ext, "engine": engine}


_WORD_FMT_MAP = {"doc": 0, "docx": 12}


def _start_word_app():
    app, engine = _dispatch_first([
        "Word.Application",       # Microsoft Word
        "kwps.Application",       # WPS 文字（常见）
        "KWPS.Application",       # WPS 文字（大小写变体）
        "wps.Application",        # 兼容别名（部分环境）
    ])
    app.Visible = False
    try:
        app.DisplayAlerts = 0
    except Exception:
        pass
    return app, engine


def _word_convert_one(app, src: Path, out: Path, target_ext: str) -> None:
    doc = app.Documents.Open(str(src), ConfirmConversions=False, ReadOnly=True, AddToRecentFiles=False)
    try:
        try:
            doc.SaveAs2(str(out), FileFormat=_WORD_FMT_MAP[target_ext])
        except Exception:
            doc.SaveAs(str(out), FileFormat=_WORD_FMT_MAP[target_ext])
    finally:
        doc.Close(SaveChanges=False)


def run_word_convert_36(
    files: list[Path], target_format: str, log_dir: Path
) -> dict[str, int]:
    log = get_logger("3_6_word_convert", log_dir)
    ext = _norm(target_format).lower().lstrip(".") or "docx"
    if ext not in _WORD_FMT_MAP:
        ext = "docx"

    valid: list[tuple[Path, Path]] = []
    ok = skip = 0
    for p in files:
        if not p.exists() or p.suffix.lower() not in (".doc", ".docx"):
            skip += 1
            log.warning("skip file=%s reason=not_word_doc", p)
            continue
        out_dir = p.parent / ext
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = out_dir / f"{p.stem}.{ext}"
        if out_path.resolve() == p.resolve():
            skip += 1
            log.info("skip file=%s reason=source_equals_target", p.name)
            continue
        if out_path.exists():
            out_path = _next_path(out_path)
        valid.append((p, out_path))

    if not valid:
        return {"ok": 0, "skip": skip, "target_ext": ext, "engine": ""}

    app, engine = _start_word_app()
    try:
        for src, out_path in valid:
            try:
                _word_convert_one(app, src, out_path, ext)
                ok += 1
                log.info("ok src=%s out=%s engine=%s", src, out_path, engine)
            except Exception as e:
                skip += 1
                log.warning("skip file=%s reason=convert_failed err=%s", src, e)
    finally:
        try:
            app.Quit()
        except Exception:
            pass
    return {"ok": ok, "skip": skip, "target_ext": ext, "engine": engine}


def _load_sheet_rename_map(cfg_path: Path) -> dict[str, str]:
    df = _load_config_rename_df(cfg_path)
    mp: dict[str, str] = {}
    for _, row in df.iterrows():
        old_name = _norm(row.get("原表名"))
        new_name = _norm(row.get("新表名"))
        if old_name and new_name:
            mp[old_name] = new_name
    return mp


def _load_file_local_rename_map(cfg_path: Path) -> list[tuple[str, str]]:
    df = _load_config_rename_df(cfg_path)
    out: list[tuple[str, str]] = []
    for _, row in df.iterrows():
        src = _norm(row.get("原文件名片段"))
        dst = _norm(row.get("新文件名片段"))
        if src:
            out.append((src, dst))
    return out


def _apply_local_name_map(name: str, mapping: list[tuple[str, str]]) -> str:
    stem = Path(name).stem
    suffix = Path(name).suffix
    new_stem = stem
    for src, dst in mapping:
        new_stem = new_stem.replace(src, dst)
    return f"{new_stem}{suffix}"


def _load_content_rename_config(cfg_path: Path) -> dict[str, object]:
    wb = load_workbook(cfg_path, read_only=True, data_only=True)
    try:
        if SHEET_CONFIG_RENAME not in wb.sheetnames:
            raise ValueError("重命名配置 缺少工作表。")
        ws = wb[SHEET_CONFIG_RENAME]
        enabled: list[dict[str, object]] = []
        for row in range(2, ws.max_row + 1):
            if not _is_enabled(ws.cell(row, 15).value):
                continue
            spec = _norm(ws.cell(row, 16).value)
            remark = _norm(ws.cell(row, 17).value)
            enabled.append({"row": row, "spec": spec, "remark": remark})
        if not enabled:
            raise ValueError("没有启用的内容重命名配置：请在 重命名配置!O:P:Q 中启用一条配置。")
        if len(enabled) > 1:
            items = []
            for item in enabled:
                remark = str(item.get("remark") or "")
                suffix = f"（{remark}）" if remark else ""
                items.append(f"第{item['row']}行{suffix}")
            raise ValueError("内容重命名配置只能启用一条：" + "、".join(items))
        return enabled[0]
    finally:
        wb.close()


def _load_content_rename_spec(cfg_path: Path) -> str:
    return str(_load_content_rename_config(cfg_path).get("spec") or "")


def _sanitize_name_part(text: str) -> str:
    s = _norm(text)
    if not s:
        return ""
    s = re.sub(r'[<>:"/\\|?*]', "_", s)
    s = re.sub(r"\s+", " ", s).strip().strip(".")
    return s


def _content_prefix_items(vals: list[str]) -> list[str]:
    return [x for x in (_sanitize_name_part(v) for v in vals) if x]


def _already_renamed_by_content(file_name: str, prefix_items: list[str]) -> bool:
    if not prefix_items:
        return False
    stem = Path(file_name).stem
    first_part = stem.split("_", 1)[0]
    haystacks = (first_part, stem)
    return all(any(item in h for h in haystacks) for item in prefix_items)


def _content_prefix(prefix_items: list[str]) -> str:
    prefix = "_".join(prefix_items)
    return f"{prefix}_" if prefix else ""


def _remove_content_prefix(file_name: str, prefix_items: list[str]) -> str | None:
    prefix = _content_prefix(prefix_items)
    if not prefix:
        return None
    stem = Path(file_name).stem
    suffix = Path(file_name).suffix
    changed = False
    while stem.startswith(prefix):
        stem = stem[len(prefix):]
        changed = True
    item_set = set(prefix_items)
    while "_" in stem:
        head, rest = stem.split("_", 1)
        if head not in item_set:
            break
        stem = rest
        changed = True
    if not changed or not stem:
        return None
    return f"{stem}{suffix}"


def _parse_sheet_cell_spec(spec: str) -> list[tuple[str, str]]:
    parts = [x.strip() for x in spec.replace("；", ";").split(";") if x and x.strip()]
    out: list[tuple[str, str]] = []
    for p in parts:
        if "@" not in p:
            raise ValueError(f"内容重命名取值规则格式错误: {p}（应为 sheet@A1）")
        sheet_name, addr = p.split("@", 1)
        sheet_name = _norm(sheet_name)
        addr = _norm(addr).upper()
        if not sheet_name or not addr:
            raise ValueError(f"内容重命名取值规则格式错误: {p}（sheet 或地址为空）")
        try:
            coordinate_from_string(addr)
        except Exception:
            raise ValueError(f"内容重命名取值规则地址非法: {p}") from None
        out.append((sheet_name, addr))
    if not out:
        raise ValueError("内容重命名取值规则为空，无法执行根据文件内容重命名。")
    return out


def _read_cells_openpyxl(path: Path, items: list[tuple[str, str]]) -> list[str]:
    wb = load_workbook(path, data_only=True, keep_vba=(path.suffix.lower() == ".xlsm"))
    try:
        vals: list[str] = []
        for sheet_name, addr in items:
            if sheet_name not in wb.sheetnames:
                vals.append(f"{sheet_name}!{addr}")
                continue
            ws = wb[sheet_name]
            v = ws[addr].value
            txt = _sanitize_name_part("" if v is None else str(v))
            vals.append(txt or f"{sheet_name}!{addr}")
        return vals
    finally:
        wb.close()
        del wb
        gc.collect()


def _read_cells_com(app, path: Path, items: list[tuple[str, str]]) -> list[str]:
    wb = app.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=True)
    try:
        vals: list[str] = []
        for sheet_name, addr in items:
            try:
                ws = wb.Worksheets(sheet_name)
            except Exception:
                vals.append(f"{sheet_name}!{addr}")
                continue
            try:
                v = ws.Range(addr).Value2
            except Exception:
                v = None
            txt = _sanitize_name_part("" if v is None else str(v))
            vals.append(txt or f"{sheet_name}!{addr}")
        return vals
    finally:
        wb.Close(SaveChanges=False)


def run_batch_rename_sheet_37(cfg_path: Path, files: list[Path], log_dir: Path) -> dict[str, int]:
    log = get_logger("3_7_batch_rename_sheet", log_dir)
    mapping = _load_sheet_rename_map(cfg_path)
    if not mapping:
        raise ValueError("重命名配置 缺少原表名/新表名映射。")

    wb_count = rename_ok = skip = 0
    for p in files:
        if not p.exists() or not _workbook_supported(p):
            skip += 1
            log.warning("skip file=%s reason=not_supported_or_missing", p)
            continue
        wb = load_workbook(p, keep_vba=(p.suffix.lower() == ".xlsm"))
        changed = False
        try:
            wb_count += 1
            for ws in wb.worksheets:
                old_name = ws.title
                lookup_name = old_name if old_name in mapping else old_name.strip()
                if lookup_name not in mapping:
                    continue
                new_name = mapping[lookup_name]
                if old_name == new_name:
                    skip += 1
                    continue
                if re.search(r"[:\\/?*\[\]]", new_name) or len(new_name) > 31 or (new_name in wb.sheetnames):
                    skip += 1
                    log.warning("skip wb=%s sheet=%s reason=invalid_or_duplicate_new_name", p.name, old_name)
                    continue
                ws.title = new_name
                rename_ok += 1
                changed = True
                log.info("ok wb=%s old=%s new=%s", p.name, old_name, new_name)
            if changed:
                wb.save(p)
        finally:
            wb.close()
    return {"workbooks": wb_count, "rename_ok": rename_ok, "skip": skip}


def run_batch_rename_by_content_46(cfg_path: Path, files: list[Path], log_dir: Path) -> dict[str, int]:
    log = get_logger("4_6_batch_rename_by_content", log_dir)
    spec_cfg = _load_content_rename_config(cfg_path)
    spec_text = str(spec_cfg.get("spec") or "")
    items = _parse_sheet_cell_spec(spec_text)

    ok = skip = 0
    com_candidates: list[Path] = []
    for p in files:
        if not p.exists() or not _excel_like(p):
            skip += 1
            log.warning("skip file=%s reason=not_excel_like_or_missing", p)
            continue
        if p.suffix.lower() in (".xls", ".xlt"):
            com_candidates.append(p)
            continue
        try:
            vals = _read_cells_openpyxl(p, items)
            prefix_items = _content_prefix_items(vals)
            if _already_renamed_by_content(p.name, prefix_items):
                skip += 1
                log.info(
                    "skip file=%s reason=already_renamed_by_content items=%s",
                    p.name, "|".join(prefix_items),
                )
                continue
            prefix = "_".join(prefix_items)[:180]
            new_name = f"{prefix}_{p.name}" if prefix else p.name
            tgt = p.with_name(new_name)
            if tgt.resolve() == p.resolve():
                skip += 1
                log.info("skip file=%s reason=name_same", p.name)
                continue
            if tgt.exists():
                tgt = _next_path(tgt)
            os.rename(p, tgt)
            ok += 1
            log.info("ok old=%s new=%s", p.name, tgt.name)
        except Exception as e:
            skip += 1
            log.warning("skip file=%s reason=rename_failed err=%s", p.name, e)

    if com_candidates:
        app = None
        try:
            app, engine = _start_excel_app()
            for p in com_candidates:
                try:
                    vals = _read_cells_com(app, p, items)
                    prefix_items = _content_prefix_items(vals)
                    if _already_renamed_by_content(p.name, prefix_items):
                        skip += 1
                        log.info(
                            "skip file=%s reason=already_renamed_by_content items=%s engine=%s",
                            p.name, "|".join(prefix_items), engine,
                        )
                        continue
                    prefix = "_".join(prefix_items)[:180]
                    new_name = f"{prefix}_{p.name}" if prefix else p.name
                    tgt = p.with_name(new_name)
                    if tgt.resolve() == p.resolve():
                        skip += 1
                        log.info("skip file=%s reason=name_same", p.name)
                        continue
                    if tgt.exists():
                        tgt = _next_path(tgt)
                    os.rename(p, tgt)
                    ok += 1
                    log.info("ok old=%s new=%s engine=%s", p.name, tgt.name, engine)
                except Exception as e:
                    skip += 1
                    log.warning("skip file=%s reason=com_rename_failed err=%s", p.name, e)
        except Exception as e:
            skip += len(com_candidates)
            log.warning("skip com_files=%s reason=start_excel_failed err=%s", len(com_candidates), e)
        finally:
            if app is not None:
                try:
                    app.Quit()
                except Exception:
                    pass
    return {
        "ok": ok,
        "skip": skip,
        "spec": spec_text,
        "config_row": int(spec_cfg.get("row") or 0),
        "config_remark": str(spec_cfg.get("remark") or ""),
    }


def run_batch_unrename_by_content_47(cfg_path: Path, files: list[Path], log_dir: Path) -> dict[str, int]:
    log = get_logger("4_7_batch_unrename_by_content", log_dir)
    spec_cfg = _load_content_rename_config(cfg_path)
    spec_text = str(spec_cfg.get("spec") or "")
    items = _parse_sheet_cell_spec(spec_text)

    ok = skip = 0
    com_candidates: list[Path] = []
    for p in files:
        if not p.exists() or not _excel_like(p):
            skip += 1
            log.warning("skip file=%s reason=not_excel_like_or_missing", p)
            continue
        if p.suffix.lower() in (".xls", ".xlt"):
            com_candidates.append(p)
            continue
        try:
            vals = _read_cells_openpyxl(p, items)
            prefix_items = _content_prefix_items(vals)
            new_name = _remove_content_prefix(p.name, prefix_items)
            if not new_name:
                skip += 1
                log.info(
                    "skip file=%s reason=prefix_not_matched items=%s",
                    p.name, "|".join(prefix_items),
                )
                continue
            tgt = p.with_name(new_name)
            if tgt.exists():
                tgt = _next_path(tgt)
            os.rename(p, tgt)
            ok += 1
            log.info("ok old=%s new=%s", p.name, tgt.name)
        except Exception as e:
            skip += 1
            log.warning("skip file=%s reason=unrename_failed err=%s", p.name, e)

    if com_candidates:
        app = None
        try:
            app, engine = _start_excel_app()
            for p in com_candidates:
                try:
                    vals = _read_cells_com(app, p, items)
                    prefix_items = _content_prefix_items(vals)
                    new_name = _remove_content_prefix(p.name, prefix_items)
                    if not new_name:
                        skip += 1
                        log.info(
                            "skip file=%s reason=prefix_not_matched items=%s engine=%s",
                            p.name, "|".join(prefix_items), engine,
                        )
                        continue
                    tgt = p.with_name(new_name)
                    if tgt.exists():
                        tgt = _next_path(tgt)
                    os.rename(p, tgt)
                    ok += 1
                    log.info("ok old=%s new=%s engine=%s", p.name, tgt.name, engine)
                except Exception as e:
                    skip += 1
                    log.warning("skip file=%s reason=com_unrename_failed err=%s", p.name, e)
        except Exception as e:
            skip += len(com_candidates)
            log.warning("skip com_files=%s reason=start_excel_failed err=%s", len(com_candidates), e)
        finally:
            if app is not None:
                try:
                    app.Quit()
                except Exception:
                    pass
    return {
        "ok": ok,
        "skip": skip,
        "spec": spec_text,
        "config_row": int(spec_cfg.get("row") or 0),
        "config_remark": str(spec_cfg.get("remark") or ""),
    }


def run_batch_rename_local_map_48(cfg_path: Path, files: list[Path], log_dir: Path) -> dict[str, int]:
    log = get_logger("4_8_batch_rename_local_map", log_dir)
    mapping = _load_file_local_rename_map(cfg_path)
    if not mapping:
        raise ValueError("重命名配置 缺少 M/N 列局部文件名映射。")

    ok = skip = 0
    for p in files:
        if not p.exists() or not p.is_file():
            skip += 1
            log.warning("skip file=%s reason=not_file_or_missing", p)
            continue
        try:
            new_name = _apply_local_name_map(p.name, mapping)
            if new_name == p.name:
                skip += 1
                log.info("skip file=%s reason=no_fragment_matched", p.name)
                continue
            tgt = p.with_name(new_name)
            if tgt.exists():
                tgt = _next_path(tgt)
            os.rename(p, tgt)
            ok += 1
            log.info("ok old=%s new=%s", p.name, tgt.name)
        except Exception as e:
            skip += 1
            log.warning("skip file=%s reason=local_map_rename_failed err=%s", p.name, e)
    return {"ok": ok, "skip": skip, "rules": len(mapping)}


def _sheet_name_invalid(new_name: str, existing_names: set[str], old_name: str) -> bool:
    if old_name == new_name:
        return True
    if re.search(r"[:\\/?*\[\]]", new_name):
        return True
    if len(new_name) > 31:
        return True
    if new_name in existing_names:
        return True
    return False


def run_batch_rename_sheet_37_com(cfg_path: Path, files: list[Path], log_dir: Path) -> dict[str, int]:
    log = get_logger("3_7_batch_rename_sheet_com", log_dir)
    mapping = _load_sheet_rename_map(cfg_path)
    if not mapping:
        raise ValueError("重命名配置 缺少原表名/新表名映射。")

    valid_files: list[Path] = []
    skip = 0
    for p in files:
        if not p.exists() or not _excel_like(p):
            skip += 1
            log.warning("skip file=%s reason=not_excel_like_or_missing", p)
            continue
        valid_files.append(p)

    if not valid_files:
        return {"workbooks": 0, "rename_ok": 0, "skip": skip, "engine": ""}

    app, engine = _start_excel_app()
    wb_count = 0
    rename_ok = 0
    try:
        for p in valid_files:
            wb = None
            changed = False
            try:
                wb = app.Workbooks.Open(str(p), UpdateLinks=0, ReadOnly=False)
                wb_count += 1
                sheet_count = wb.Worksheets.Count
                existing_names = set()
                for idx in range(1, sheet_count + 1):
                    existing_names.add(str(wb.Worksheets(idx).Name))

                for idx in range(1, sheet_count + 1):
                    ws = wb.Worksheets(idx)
                    old_name = str(ws.Name)
                    lookup_name = old_name if old_name in mapping else old_name.strip()
                    if lookup_name not in mapping:
                        continue
                    new_name = mapping[lookup_name]
                    names_without_old = {x for x in existing_names if x != old_name}
                    if _sheet_name_invalid(new_name, names_without_old, old_name):
                        skip += 1
                        log.warning(
                            "skip wb=%s sheet=%s reason=invalid_or_duplicate_new_name",
                            p.name, old_name
                        )
                        continue
                    ws.Name = new_name
                    existing_names.discard(old_name)
                    existing_names.add(new_name)
                    rename_ok += 1
                    changed = True
                    log.info("ok wb=%s old=%s new=%s", p.name, old_name, new_name)

                if changed:
                    wb.Save()
            except Exception as e:
                skip += 1
                log.warning("skip wb=%s reason=com_rename_failed err=%s", p.name, e)
            finally:
                if wb is not None:
                    try:
                        wb.Close(SaveChanges=False)
                    except Exception:
                        pass
    finally:
        try:
            app.Quit()
        except Exception:
            pass

    return {"workbooks": wb_count, "rename_ok": rename_ok, "skip": skip, "engine": engine}
